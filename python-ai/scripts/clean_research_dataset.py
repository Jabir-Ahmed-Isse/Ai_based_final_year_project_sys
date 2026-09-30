"""Clean the Hormuud research workbook into reproducible JSON and CSV inputs."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

WORKSHEET = "Project Test Dataset"
HEADERS = [
    "No.",
    "Project Title",
    "Description",
    "Problem Statement",
    "Research Objectives",
    "Features",
    "Technologies and Tools",
]
FIELD_MAP = {
    "No.": "number",
    "Project Title": "title",
    "Description": "description",
    "Problem Statement": "problemStatement",
    "Research Objectives": "objectives",
    "Features": "features",
    "Technologies and Tools": "technologiesAndTools",
}


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFC", str(value))
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\s*\n+\s*", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def normalized_title(value: str) -> str:
    text = unicodedata.normalize("NFKC", value).casefold()
    text = re.sub(r"[^\w]+", " ", text, flags=re.UNICODE)
    return re.sub(r"\s+", " ", text).strip()


def content_hash(row: dict[str, Any]) -> str:
    canonical = "\u241f".join(clean_text(row[field]).casefold() for field in (
        "title", "description", "problemStatement", "objectives", "features", "technologiesAndTools"
    ))
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def clean_workbook(input_path: Path) -> dict[str, Any]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    if WORKSHEET not in workbook.sheetnames:
        raise ValueError(f"Required worksheet '{WORKSHEET}' was not found")
    worksheet = workbook[WORKSHEET]
    actual_headers = [clean_text(cell.value) for cell in worksheet[1]][: len(HEADERS)]
    if actual_headers != HEADERS:
        raise ValueError(f"Unexpected headers: {actual_headers}")

    projects: list[dict[str, Any]] = []
    invalid_rows: list[dict[str, Any]] = []
    duplicates: list[dict[str, Any]] = []
    seen_titles: dict[str, int] = {}
    seen_hashes: dict[str, int] = {}
    total_rows = 0
    empty_rows = 0

    for source_row, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        raw = dict(zip(HEADERS, values[: len(HEADERS)]))
        if not any(clean_text(value) for value in raw.values()):
            empty_rows += 1
            continue
        total_rows += 1
        project = {FIELD_MAP[header]: clean_text(raw.get(header)) for header in HEADERS}
        if project["number"]:
            try:
                project["number"] = int(float(project["number"]))
            except ValueError:
                pass
        project = {"sourceRow": source_row, **project}
        if not project["title"]:
            invalid_rows.append({"sourceRow": source_row, "reason": "Missing project title", "row": project})
            continue
        project["normalizedTitle"] = normalized_title(project["title"])
        project["contentHash"] = content_hash(project)
        duplicate_of = seen_titles.get(project["normalizedTitle"]) or seen_hashes.get(project["contentHash"])
        if duplicate_of:
            duplicates.append({"sourceRow": source_row, "duplicateOfSourceRow": duplicate_of, "row": project})
            continue
        seen_titles[project["normalizedTitle"]] = source_row
        seen_hashes[project["contentHash"]] = source_row
        projects.append(project)

    workbook.close()
    return {
        "sourceFile": input_path.name,
        "sourceWorksheet": WORKSHEET,
        "notesExcluded": True,
        "syntheticContent": True,
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "projects": projects,
        "duplicates": duplicates,
        "invalidRows": invalid_rows,
        "summary": {
            "rowsRead": total_rows,
            "projectsAccepted": len(projects),
            "duplicatesRemoved": len(duplicates),
            "invalidRows": len(invalid_rows),
            "emptyRowsRemoved": empty_rows,
        },
        "warnings": [
            "The Notes worksheet was deliberately excluded.",
            "Descriptive fields in the supplied workbook are synthetic test content inferred from project titles.",
        ],
    }


def write_outputs(result: dict[str, Any], json_path: Path, csv_path: Path, report_path: Path) -> None:
    for target in (json_path, csv_path, report_path):
        target.parent.mkdir(parents=True, exist_ok=True)
    json_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    fields = [
        "sourceRow", "number", "title", "description", "problemStatement", "objectives",
        "features", "technologiesAndTools", "normalizedTitle", "contentHash",
    ]
    with csv_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(result["projects"])
    report_path.write_text(json.dumps({
        **result["summary"],
        "sourceFile": result["sourceFile"],
        "sourceWorksheet": result["sourceWorksheet"],
        "notesExcluded": result["notesExcluded"],
        "warnings": result["warnings"],
        "invalidRows": result["invalidRows"],
        "duplicates": result["duplicates"],
    }, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output-json", required=True, type=Path)
    parser.add_argument("--output-csv", required=True, type=Path)
    parser.add_argument("--report", required=True, type=Path)
    args = parser.parse_args()
    result = clean_workbook(args.input)
    write_outputs(result, args.output_json, args.output_csv, args.report)
    print(json.dumps(result["summary"]))


if __name__ == "__main__":
    main()
