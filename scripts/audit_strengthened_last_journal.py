from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader


def metric_check(tp: int, tn: int, fp: int, fn: int, expected: list[float]) -> dict:
    n = tp + tn + fp + fn
    accuracy = 100 * (tp + tn) / n
    precision = 100 * tp / (tp + fp) if tp + fp else 0
    recall = 100 * tp / (tp + fn) if tp + fn else 0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0
    actual = [accuracy, precision, recall, f1]
    return {
        "n": n,
        "actual": [round(x, 2) for x in actual],
        "expected": expected,
        "passed": all(math.isclose(a, e, abs_tol=0.011) for a, e in zip(actual, expected)),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("pdf", type=Path)
    parser.add_argument("change_report", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    doc = Document(args.docx)
    change = json.loads(args.change_report.read_text(encoding="utf-8"))
    refs = [p.text for p in doc.paragraphs if re.match(r"^\[\d+\]", p.text.strip())]
    ref_numbers = [int(re.match(r"^\[(\d+)\]", text.strip()).group(1)) for text in refs]
    text = "\n".join(p.text for p in doc.paragraphs)
    all_tables = [[cell.text.strip() for row in table.rows for cell in row.cells] for table in doc.tables]

    similarity = {
        "TF-IDF": metric_check(1, 19, 0, 10, [66.67, 100.00, 9.09, 16.67]),
        "Sentence-BERT": metric_check(12, 16, 2, 0, [93.33, 85.71, 100.00, 92.31]),
        "BGE-M3": metric_check(27, 0, 3, 0, [90.00, 90.00, 100.00, 94.74]),
    }
    supervisor = {
        "TF-IDF": metric_check(206, 1262, 63, 29, [94.10, 76.58, 87.66, 81.75]),
        "Sentence-BERT": metric_check(150, 1296, 30, 84, [92.69, 83.33, 64.10, 72.46]),
        "BGE-M3": metric_check(177, 1261, 65, 57, [92.18, 73.14, 75.64, 74.37]),
    }

    with ZipFile(args.docx) as archive:
        bad_zip = archive.testzip()
    pdf = PdfReader(str(args.pdf))
    blank_pages = [i + 1 for i, p in enumerate(pdf.pages) if not (p.extract_text() or "").strip()]

    report = {
        "docxSha256": hashlib.sha256(args.docx.read_bytes()).hexdigest(),
        "zipCrcFailure": bad_zip,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "inlineShapes": len(doc.inline_shapes),
        "referenceCount": len(refs),
        "referencesSequential": ref_numbers == list(range(1, 115)),
        "referencesUnchanged": change["referenceListByteEquivalent"],
        "citationSequencePreserved": change["citationSequencePreserved"],
        "similarityMetrics": similarity,
        "supervisorMetrics": supervisor,
        "allConfusionMetricsRecomputed": all(x["passed"] for x in similarity.values()) and all(x["passed"] for x in supervisor.values()),
        "rankingStatementsPresent": {
            "sentenceBertBalancedFirst": "Sentence-BERT leads the evidence-balanced similarity ranking" in text,
            "bgeOrdinaryF1First": "BGE-M3 retains the highest ordinary similarity F1" in text,
            "tfidfSupervisorFirst": "TF-IDF leads on accuracy, F1, and MCC" in text,
            "criterionSpecific": "criterion-specific" in text,
        },
        "evidenceIntegrity": {
            "commonPairedSetNotClaimed": "they are not presented as a common paired test set" in text,
            "multiRaterNotClaimed": "no multi-rater agreement or adjudicated consensus is claimed" in text,
            "externalCohortNotClaimed": "no external cohort is available" in text,
            "futureValidationNotCurrentEvidence": "not as evidence already obtained" in text,
        },
        "pdfPages": len(pdf.pages),
        "blankPdfPages": blank_pages,
    }
    report["passed"] = all([
        bad_zip is None,
        len(doc.paragraphs) == 299,
        len(doc.tables) == 24,
        len(doc.inline_shapes) == 15,
        len(refs) == 114,
        report["referencesSequential"],
        report["referencesUnchanged"],
        report["citationSequencePreserved"],
        report["allConfusionMetricsRecomputed"],
        all(report["rankingStatementsPresent"].values()),
        all(report["evidenceIntegrity"].values()),
        len(pdf.pages) == 20,
        not blank_pages,
    ])
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps(report, indent=2))
    if not report["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
