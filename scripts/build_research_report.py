"""Build the measured research report in DOCX and HTML formats."""
from __future__ import annotations

import argparse
import html
import json
import math
import statistics
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
GOLD = "BF9000"
INK = "1F2937"
MUTED = "667085"
GREEN = "E2F0D9"
RED = "9B1C1C"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120

MODEL_ORDER = ["tfidf", "sentence_bert", "bge_m3"]
MODEL_LABELS = {"tfidf": "TF-IDF", "sentence_bert": "Sentence-BERT", "bge_m3": "BGE-M3"}


def mean(values: list[float]) -> float:
    return statistics.fmean(values) if values else 0.0


def percentile(values: list[float], p: float) -> float:
    values = sorted(values)
    if not values:
        return 0.0
    index = (len(values) - 1) * p
    lower, upper = math.floor(index), math.ceil(index)
    if lower == upper:
        return values[lower]
    return values[lower] * (upper - index) + values[upper] * (index - lower)


def gini(values: list[int]) -> float:
    values = sorted(max(0, value) for value in values)
    if not values or sum(values) == 0:
        return 0.0
    total = sum((index + 1) * value for index, value in enumerate(values))
    return (2 * total) / (len(values) * sum(values)) - (len(values) + 1) / len(values)


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def set_cell_shading(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    properties = table._tbl.tblPr
    width = properties.first_child_found_in("w:tblW")
    if width is None:
        width = OxmlElement("w:tblW")
        properties.append(width)
    width.set(qn("w:w"), str(sum(widths)))
    width.set(qn("w:type"), "dxa")
    indent = properties.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    indent.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(value))
        grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            properties = cell._tc.get_or_add_tcPr()
            cell_width = properties.first_child_found_in("w:tcW")
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                properties.append(cell_width)
            cell_width.set(qn("w:w"), str(widths[index]))
            cell_width.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def set_font(run, size: float | None = None, color: str | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_paragraph(doc: Document, text: str = "", *, bold_lead: str | None = None, style: str | None = None, align=None, color: str = INK) -> Any:
    paragraph = doc.add_paragraph(style=style)
    if align is not None:
        paragraph.alignment = align
    if bold_lead and text.startswith(bold_lead):
        first = paragraph.add_run(bold_lead)
        set_font(first, bold=True, color=color)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_font(rest, color=color)
    else:
        run = paragraph.add_run(text)
        set_font(run, color=color)
    return paragraph


def add_heading(doc: Document, text: str, level: int = 1) -> Any:
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_table(doc: Document, headers: list[str], rows: list[list[Any]], widths: list[int], font_size: float = 8.5) -> Any:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    repeat_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        set_font(run, size=font_size, color="FFFFFF", bold=True)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            if row_index % 2:
                set_cell_shading(cell, "F8FAFC")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if isinstance(value, (int, float)) else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(str(value))
            set_font(run, size=font_size, color=INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_font(run, size=9, color=MUTED, italic=True)


def add_figure(doc: Document, image: Path, caption: str, width: float = 6.2) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    shape = run.add_picture(str(image), width=Inches(width))
    try:
        shape._inline.docPr.set("descr", caption)
    except Exception:
        pass
    add_caption(doc, caption)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    # Narrative-proposal list tokens. The built-in list styles preserve true
    # Word numbering while these paragraph settings keep the manuscript compact.
    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor.from_string(INK)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("Comparative AI Model Evaluation | Hormuud CSIT")
    set_font(run, size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Research report | ")
    set_font(run, size=8.5, color=MUTED)
    add_page_field(footer)


def build_data(raw_dir: Path) -> dict[str, Any]:
    projects = load_json(raw_dir / "projects.json")
    users = load_json(raw_dir / "test-users.json")
    supervisors = load_json(raw_dir / "supervisor-profiles.json")
    experiments = load_json(raw_dir / "experiment-runs.json")
    pairs = load_json(raw_dir / "project-pair-scores.json")
    supervisor_scores = load_json(raw_dir / "supervisor-matching-scores.json")
    assignments = load_json(raw_dir / "balanced-assignments.json")
    imports = load_json(raw_dir / "import-runs.json")
    activities = load_json(raw_dir / "activity-logs.json")

    model_stats = []
    for model in MODEL_ORDER:
        rows = [row for row in pairs if row["model"] == model]
        scores = [float(row["weightedOverallScore"]) for row in rows]
        times = [float(row["executionTimeMs"]) for row in rows]
        risks = Counter(row["riskLevel"] for row in rows)
        project_run = next((run for run in experiments if run["type"] == "project_similarity" and run["status"] == "completed" and run["models"] == [model]), {})
        hardware = (project_run.get("hardware") or {}).get(model, {})
        supervisor_run = next((run for run in experiments if run["type"] == "supervisor_matching" and run["status"] == "completed" and run["models"] == [model]), {})
        model_stats.append({
            "key": model,
            "label": MODEL_LABELS[model],
            "count": len(rows),
            "mean": mean(scores),
            "median": percentile(scores, 0.5),
            "std": statistics.pstdev(scores),
            "min": min(scores),
            "max": max(scores),
            "low": risks["Low Risk"],
            "medium": risks["Medium Risk"],
            "high": risks["High Risk"],
            "cached_pair_ms": mean(times),
            "loading_s": float(hardware.get("loading_time_ms") or 0) / 1000,
            "embedding_dimension": hardware.get("embedding_dimension") or "n/a",
            "project_duration_s": float(project_run.get("durationMs") or 0) / 1000,
            "supervisor_duration_s": float(supervisor_run.get("durationMs") or 0) / 1000,
        })

    workloads = Counter(row.get("assignedSupervisorCode") for row in assignments)
    load_values = [workloads.get(row["supervisorCode"], 0) for row in supervisors]
    capacities = {row["supervisorCode"]: int(row["maximumCapacity"]) for row in supervisors}
    capacity_exceeded = sum(1 for code, load in workloads.items() if load > capacities[code])
    assignment_stats = {
        "supervisors_used": sum(1 for value in load_values if value > 0),
        "mean": mean(load_values),
        "std": statistics.pstdev(load_values),
        "variance": statistics.pvariance(load_values),
        "gini": gini(load_values),
        "min": min(load_values),
        "max": max(load_values),
        "capacity_exceeded": capacity_exceeded,
    }

    pair_map: dict[str, dict[str, dict]] = defaultdict(dict)
    for row in pairs:
        key = ":".join(sorted([str(row["firstProjectId"]), str(row["secondProjectId"])]))
        pair_map[key][row["model"]] = row
    disagreements = []
    for model_rows in pair_map.values():
        if len(model_rows) != 3:
            continue
        values = [float(model_rows[model]["weightedOverallScore"]) for model in MODEL_ORDER]
        sample = model_rows["tfidf"]
        disagreements.append({
            "first": sample["firstProjectTitle"],
            "second": sample["secondProjectTitle"],
            "tfidf": values[0],
            "sentence": values[1],
            "bge": values[2],
            "spread": max(values) - min(values),
        })
    disagreements.sort(key=lambda row: row["spread"], reverse=True)
    return {
        "projects": projects,
        "users": users,
        "supervisors": supervisors,
        "experiments": experiments,
        "pairs": pairs,
        "supervisor_scores": supervisor_scores,
        "assignments": assignments,
        "imports": imports,
        "activities": activities,
        "model_stats": model_stats,
        "assignment_stats": assignment_stats,
        "disagreements": disagreements[:20],
    }


def build_docx(data: dict[str, Any], charts_dir: Path, output: Path) -> None:
    doc = Document()
    configure_document(doc)

    # editorial_cover override for the opening page.
    for _ in range(5):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("ACADEMIC RESEARCH REPORT")
    set_font(run, size=10, color=GOLD, bold=True)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    run = title.add_run("Comparative Evaluation of TF-IDF, Sentence-BERT and BGE-M3")
    set_font(run, size=27, color=NAVY, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Academic Project Similarity Detection and Supervisor Assignment")
    set_font(run, size=15, color=DARK_BLUE)
    metadata = doc.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_before = Pt(50)
    run = metadata.add_run("Hormuud University | Faculty of Computer Science and IT\nMeasured experiment: 24 July 2026 | Random seed: 42")
    set_font(run, size=11, color=MUTED)
    warning = doc.add_paragraph()
    warning.alignment = WD_ALIGN_PARAGRAPH.CENTER
    warning.paragraph_format.space_before = Pt(28)
    run = warning.add_run("RESEARCH STATUS: DESCRIPTIVE RESULTS COMPLETE; HUMAN GROUND-TRUTH VALIDATION PENDING")
    set_font(run, size=10, color=RED, bold=True)
    doc.add_page_break()

    add_heading(doc, "Contents", 1)
    for item in [
        "1-8. Context, dataset and experimental design",
        "9-15. Ground truth, annotation and evaluation protocol",
        "16-24. Measured results, efficiency, error analysis and workload",
        "25-30. Validity, ethics, reproducibility, deployment and conclusion",
        "31-32. References and appendices",
    ]:
        add_paragraph(doc, item)
    doc.add_page_break()

    add_heading(doc, "1. Abstract", 1)
    add_paragraph(doc, "This report documents a complete local research experiment using 78 imported Computer Science and IT final-year projects, 20 structured supervisor profiles, and three similarity model families: TF-IDF, multilingual Sentence-BERT, and BGE-M3. The system generated and stored all 3,003 unique project pairs per model (9,009 model-pair records) and all 1,560 project-supervisor combinations per model (4,680 matching records). A capacity-aware ensemble procedure produced 78 provisional assignments using all 20 supervisors with no capacity violations. The source workbook contains no expert ground truth. Consequently, score distributions, correlations, field scores, timings and workload measures are reported, while accuracy, F1, ROC-AUC, PR-AUC, MRR, MAP, NDCG and best-model claims are withheld. A separate 600-pair synthetic construction benchmark is supplied only to prepare the validation workflow; it is not presented as human-validated evidence.")

    add_heading(doc, "2. Introduction", 1)
    add_paragraph(doc, "Final-year project lifecycle systems must detect idea overlap without confusing shared technologies or generic management-system language with duplicated research contributions. They must also recommend supervisors from structured academic evidence while respecting capacity. The experiment extends the existing application rather than creating an independent research system.")

    add_heading(doc, "3. Research problem", 1)
    add_paragraph(doc, "The central research problem is to determine how lexical and multilingual semantic representations behave on the same complete institutional dataset for two linked tasks: structured six-field project similarity and structured six-field supervisor expertise matching. A second problem is methodological: the supplied data do not contain expert labels, so performance claims must be separated from descriptive model behaviour.")

    add_heading(doc, "4. Research objectives", 1)
    objectives = [
        ("Objective 1.", " Import and audit all valid records without deleting unrelated application data."),
        ("Objective 2.", " Compare TF-IDF, multilingual Sentence-BERT and BGE-M3 over every unique recorded project pair."),
        ("Objective 3.", " Compare the same model families over every project-supervisor combination using structured profiles."),
        ("Objective 4.", " Generate capacity-safe provisional assignments while retaining pure semantic and workload-adjusted rankings."),
        ("Objective 5.", " Provide reproducible raw data, dashboards, annotations, visualizations and reports without fabricating label-dependent metrics."),
    ]
    for lead, text in objectives:
        add_paragraph(doc, lead + text, bold_lead=lead)

    add_heading(doc, "5. Dataset description", 1)
    add_table(doc, ["Item", "Measured value"], [
        ["Source worksheet", "Project Test Dataset"],
        ["Imported projects", "78"],
        ["Required fields", "Title, description, problem statement, objectives, features, technologies/tools"],
        ["Notes worksheet imported", "No"],
        ["Test students", "78 (one per project)"],
        ["Test supervisors", "20"],
        ["Unique project pairs", "3,003"],
        ["Project-supervisor combinations", "1,560 per model"],
        ["Workbook content warning", "Descriptive project fields are synthetic test content inferred from titles"],
    ], [2700, 6660])

    add_heading(doc, "6. Data-cleaning procedure", 1)
    import_run = data["imports"][0]
    counts = import_run.get("counts") or {}
    add_table(doc, ["Cleaning or import measure", "Result"], [
        ["Rows read", counts.get("rowsRead", 78)],
        ["Projects imported", counts.get("newlyCreatedRecords", 78)],
        ["Duplicates removed", counts.get("duplicatesRemoved", 0)],
        ["Invalid rows", counts.get("invalidRows", 0)],
        ["Skipped rows", counts.get("skippedRows", 0)],
        ["Unrelated projects preserved", counts.get("preservedProjects", 100)],
        ["Notes excluded", "Yes"],
        ["Duplicate keys", "Normalized title and SHA-256 six-field content hash"],
        ["Transaction status", "MongoDB standalone: transaction unavailable; backup-first idempotent tagged import used"],
    ], [4000, 5360])

    add_heading(doc, "7. Supervisor-profile dataset", 1)
    add_paragraph(doc, "Each test supervisor has research interests, areas of expertise, academic specialization, skills, technologies, previous supervised topics, publication/research keywords, years of experience, maximum capacity, current workload and availability. Profiles are synthetic, deterministic, tagged test data.")
    add_table(doc, ["Code", "Academic specialization", "Capacity"], [
        [row["supervisorCode"], row["academicSpecialization"], row["maximumCapacity"]]
        for row in data["supervisors"]
    ], [1200, 6360, 1800], 8)

    add_heading(doc, "8. Experimental design", 1)
    add_paragraph(doc, "The experiment uses fixed seed 42, batch embeddings, process-wide model caching and persisted run identifiers. Every project pair is generated exactly once with canonical pair direction. Each model stores field-level scores, weighted overall score, unweighted combined-text score, risk level, execution time, model version, embedding metadata, device and batch size. Supervisor matching stores all 20 scores for every project, pure semantic rank and workload-adjusted rank.")

    add_heading(doc, "9. Ground-truth creation", 1)
    add_paragraph(doc, "No expert ground truth was present. A provisional 600-pair synthetic construction benchmark was generated with 200 different, 200 partially related and 200 highly similar items. Its 60/20/20 development, validation and test partitions are grouped by source project, and automated checks confirm no project-level leakage. These construction labels are not substitutes for expert judgement and are excluded from the measured findings in this report.")

    add_heading(doc, "10. Annotation procedure", 1)
    add_paragraph(doc, "The administrator annotation interface presents two complete recorded project profiles. Independent annotators select 0 (different), 1 (partially related) or 2 (highly similar), may add a 0-100 continuous score and notes, and submit under a named dataset. Labels are unique per evaluator, pair and dataset. The interface also supports 0-3 supervisor relevance grades with multiple acceptable supervisors per project.")

    add_heading(doc, "11. Inter-annotator agreement", 1)
    add_paragraph(doc, "No human annotations existed at report time; therefore Cohen's kappa, Fleiss' kappa and Krippendorff's alpha are not computed. Consensus records are prepared and require at least two independent labels, preferably three. Disputed pairs are marked for adjudication rather than forced into consensus.")

    doc.add_page_break()
    add_heading(doc, "12. Model configurations", 1)
    add_table(doc, ["Model", "Exact model/configuration", "Embedding"], [
        ["TF-IDF", "Word unigrams/bigrams (70%) + character 3-5 grams (30%); lowercase; L2; cosine", "Sparse"],
        ["Sentence-BERT", "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2; normalized dense vectors; batch 32", "384 dimensions"],
        ["BGE-M3", "BAAI/bge-m3; normalized dense vectors; batch 32", "1,024 dimensions"],
    ], [1700, 5860, 1800])

    add_heading(doc, "13. Project-text representation", 1)
    add_table(doc, ["Field", "Configured weight"], [
        ["Project title", "20%"],
        ["Description", "20%"],
        ["Problem statement", "20%"],
        ["Research objectives", "15%"],
        ["Features", "15%"],
        ["Technologies and tools", "10%"],
    ], [6000, 3360])
    add_paragraph(doc, "Each field is scored independently. The weighted score is normalized over available fields, and a separate unweighted combined-text score is stored. The title is never used alone.")

    add_heading(doc, "14. Similarity thresholds", 1)
    add_paragraph(doc, "The descriptive thresholds are Low Risk 0-39%, Medium Risk 40-69% and High Risk 70-100%. They are configurable and have not been optimized against the final test set. Threshold optimization must use only the future expert-labelled validation partition.")

    add_heading(doc, "15. Evaluation metrics", 1)
    add_paragraph(doc, "The implementation supports binary, multiclass, continuous-score, ranking, efficiency and confidence-interval metrics. Label-dependent metrics remain pending. The current evidence is limited to score distributions, cross-model correlation, field-level scores, runtime metadata, supervisor coverage, workload distribution, capacity violations and ranking stability.")

    add_heading(doc, "16. Similarity-detection results", 1)
    add_table(doc, ["Model", "Pairs", "Mean", "Median", "Std. dev.", "Low", "Medium", "High"], [
        [row["label"], row["count"], f'{row["mean"]:.2f}', f'{row["median"]:.2f}', f'{row["std"]:.2f}', row["low"], row["medium"], row["high"]]
        for row in data["model_stats"]
    ], [1700, 900, 1050, 1050, 1050, 900, 900, 910], 8)
    add_paragraph(doc, "Score scale differences are substantial: BGE-M3 produces the highest descriptive mean, followed by Sentence-BERT and TF-IDF. This does not establish accuracy or a winner because calibration against expert labels is absent.")
    add_figure(doc, charts_dir / "figure-01-similarity-distribution.png", "Figure 1. Weighted similarity-score distributions for all 3,003 recorded pairs per model.")
    add_figure(doc, charts_dir / "figure-04-model-correlation.png", "Figure 2. Spearman correlation of model scores for the same recorded project pairs.", 5.7)

    add_heading(doc, "17. Supervisor-assignment results", 1)
    add_paragraph(doc, "All three model families stored 1,560 project-supervisor records. The operational assignment stage averaged available model scores and applied a progressive utilization penalty while enforcing each supervisor's capacity. All 78 projects received one provisional supervisor; all 20 supervisors were used; no capacity was exceeded.")
    add_figure(doc, charts_dir / "figure-08-supervisor-heatmap.png", "Figure 3. BGE-M3 adjusted project-supervisor scores for all 78 x 20 combinations.", 6.35)

    add_heading(doc, "18. Efficiency results", 1)
    add_table(doc, ["Model", "Initial load (s)", "Project run (s)", "Supervisor run (s)", "Cached batch ms/pair"], [
        [row["label"], f'{row["loading_s"]:.2f}', f'{row["project_duration_s"]:.2f}', f'{row["supervisor_duration_s"]:.2f}', f'{row["cached_pair_ms"]:.4f}']
        for row in data["model_stats"]
    ], [1900, 1800, 1800, 1900, 1960], 8.5)
    add_paragraph(doc, "BGE-M3 required the largest initial model load. The successful BGE-M3 project rerun reused the populated embedding cache, so its stored per-pair value must not be interpreted as cold-start latency. Sentence-BERT's first project run includes its initial download/load. TF-IDF is the fastest and simplest measured baseline.")
    add_figure(doc, charts_dir / "figure-05-latency.png", "Figure 4. Stored post-cache execution time per pair with 95% confidence intervals.", 5.8)

    add_heading(doc, "19. Statistical significance tests", 1)
    add_paragraph(doc, "McNemar, Wilcoxon signed-rank, paired bootstrap error tests and effect sizes require reference outcomes or errors. They are intentionally not reported. Bootstrap intervals shown for raw mean scores quantify descriptive sampling uncertainty only, not accuracy uncertainty.")
    add_figure(doc, charts_dir / "figure-11-score-confidence-intervals.png", "Figure 5. Descriptive mean-score bootstrap intervals; not model-accuracy intervals.", 5.8)

    add_heading(doc, "20. Model comparison", 1)
    add_table(doc, ["Model", "Measured operational strength", "Main limitation", "Current evidence status"], [
        ["TF-IDF", "Fast, transparent lexical baseline", "Weak paraphrase and multilingual semantics", "Complete descriptive run; quality unvalidated"],
        ["Sentence-BERT", "Multilingual semantic representation with smaller dense vectors", "Cold start and semantic calibration cost", "Complete descriptive run; quality unvalidated"],
        ["BGE-M3", "High-dimensional multilingual technical representation", "Largest cold-start and deployment footprint", "Complete descriptive run; quality unvalidated"],
    ], [1600, 3000, 2860, 1900], 8)
    add_paragraph(doc, "No quality rank is assigned. A higher unlabelled mean score can reflect calibration rather than better discrimination. The final decision matrix must be populated after expert-labelled macro F1, PR-AUC, Spearman correlation, NDCG@3 and MRR exist.")

    add_heading(doc, "21. Error analysis", 1)
    add_paragraph(doc, "False positives and false negatives cannot be identified without ground truth. The current error-analysis substitute is model disagreement analysis. It highlights pairs needing annotation, especially generic management systems, shared technologies with different purposes, semantically related purposes using different technologies, short descriptions, mixed Somali-English text and ambiguous titles.")
    add_table(doc, ["First project", "Second project", "TF-IDF", "Sentence-BERT", "BGE-M3", "Spread"], [
        [row["first"], row["second"], f'{row["tfidf"]:.1f}', f'{row["sentence"]:.1f}', f'{row["bge"]:.1f}', f'{row["spread"]:.1f}']
        for row in data["disagreements"][:8]
    ], [2250, 2250, 1050, 1300, 1050, 1460], 7.5)

    add_heading(doc, "22. Field-weight ablation study", 1)
    add_paragraph(doc, "Field-level scores are complete, but ablation performance requires labels. The current chart reports mean field scores only and must not be interpreted as causal importance. A future validation-only grid search should vary weights, freeze the selected configuration, and evaluate once on the held-out test partition.")
    add_figure(doc, charts_dir / "figure-06-field-similarity.png", "Figure 6. Descriptive mean field-level similarity by model; not an ablation result.", 6.1)

    add_heading(doc, "23. Workload-balancing evaluation", 1)
    assignment = data["assignment_stats"]
    add_table(doc, ["Measure", "Result"], [
        ["Projects assigned", "78"],
        ["Supervisors used", assignment["supervisors_used"]],
        ["Mean projects per supervisor", f'{assignment["mean"]:.2f}'],
        ["Workload standard deviation", f'{assignment["std"]:.3f}'],
        ["Workload variance", f'{assignment["variance"]:.3f}'],
        ["Gini coefficient", f'{assignment["gini"]:.3f}'],
        ["Minimum / maximum workload", f'{assignment["min"]} / {assignment["max"]}'],
        ["Capacity violations", assignment["capacity_exceeded"]],
    ], [5000, 4360])
    add_figure(doc, charts_dir / "figure-13-workload-after.png", "Figure 7. Provisional supervisor workload after capacity-aware balancing.", 6.2)

    add_heading(doc, "24. Discussion", 1)
    add_paragraph(doc, "The experiment demonstrates complete database coverage and stable integration of traditional and semantic models. TF-IDF separates low lexical overlap aggressively. Sentence-BERT and BGE-M3 produce broader semantic score distributions. The model families also vary in load cost and embedding size. These behaviours are operational facts, not evidence of superior relevance. The annotation queue should prioritize strong disagreement pairs to reduce evaluation effort.")

    add_heading(doc, "25. Threats to validity", 1)
    add_paragraph(doc, "Construct validity is limited by synthetic descriptive project fields and absent expert labels. Internal validity is limited by a single institutional dataset and CPU execution environment. External validity is limited to CSIT project topics from the supplied workbook. Score calibration differs across model families. The BGE-M3 cache-assisted rerun separates cold-start load from cached scoring and must be interpreted accordingly.")

    add_heading(doc, "26. Ethical considerations", 1)
    add_paragraph(doc, "All created accounts, projects and supervisor profiles are marked test data. The report uses synthetic names and does not expose production passwords or private personal information. Recommendations are decision support: supervisors and administrators must retain review and override authority. Capacity, specialization and availability should not become opaque barriers to equitable allocation.")

    add_heading(doc, "27. Limitations", 1)
    add_paragraph(doc, "No human-validated accuracy or ranking quality is available. The synthetic benchmark is not expert ground truth. BGE-M3 uses dense embeddings only; sparse and multi-vector modes were not activated. Memory was observed operationally but not captured with a controlled profiler for every run. Domain-level evaluation awaits a validated project-domain taxonomy.")

    add_heading(doc, "28. Reproducibility information", 1)
    add_table(doc, ["Item", "Recorded value"], [
        ["Operating environment", "Windows desktop, local services"],
        ["Python", "3.13.14"],
        ["Node.js", "24.18.0"],
        ["Database", "MongoDB hormuud-gpms"],
        ["Random seed", "42"],
        ["Project models", "TF-IDF; paraphrase-multilingual-MiniLM-L12-v2; BAAI/bge-m3"],
        ["Project pairs per model", "3,003"],
        ["Supervisor scores per model", "1,560"],
        ["Git revision", "Unavailable: workspace contains no usable commit metadata"],
        ["Configuration", "config/research-experiment.json"],
        ["Dependency locks", "python-ai/requirements.txt; backend and frontend package-lock.json"],
    ], [3000, 6360])

    add_heading(doc, "29. Recommended deployment model", 1)
    add_paragraph(doc, "No final quality winner can be selected. For staged deployment, retain TF-IDF as an explainable baseline, run multilingual Sentence-BERT in shadow mode for semantic decision support, and keep BGE-M3 as an offline research candidate until expert validation demonstrates a justified quality gain relative to its cold-start and resource costs. Revisit this recommendation only after consensus labels and held-out evaluation.")

    add_heading(doc, "30. Conclusion", 1)
    add_paragraph(doc, "The implementation completed the full unlabelled experiment safely and reproducibly: 78 projects, 9,009 pair scores, 4,680 supervisor scores and 78 capacity-safe assignments. It also provides the annotation, consensus, export, visualization and audit foundations required for a valid second phase. The responsible conclusion is that all three models are operational, their behaviours differ materially, and model quality remains undetermined until expert ground truth is collected.")

    add_heading(doc, "31. References", 1)
    references = [
        "Manning, C. D., Raghavan, P., and Schuetze, H. Introduction to Information Retrieval. Cambridge University Press, 2008.",
        "Reimers, N., and Gurevych, I. Sentence-BERT: Sentence Embeddings using Siamese BERT-Networks. EMNLP-IJCNLP, 2019.",
        "Chen, J. et al. BGE M3-Embedding: Multi-Lingual, Multi-Functionality, Multi-Granularity Text Embeddings. arXiv, 2024.",
        "scikit-learn documentation: TfidfVectorizer and cosine similarity.",
        "Sentence-Transformers documentation: semantic textual similarity and multilingual models.",
    ]
    for index, reference in enumerate(references, 1):
        add_paragraph(doc, f"{index}. {reference}")

    add_heading(doc, "32. Appendices", 1)
    add_heading(doc, "Appendix A. Run identifiers", 2)
    add_table(doc, ["Run ID", "Type", "Model(s)", "Status", "Duration (s)"], [
        [row["experimentRunId"], row["type"], ", ".join(row["models"]), row["status"], f'{float(row.get("durationMs") or 0)/1000:.2f}']
        for row in data["experiments"]
    ], [2600, 1900, 2100, 1200, 1560], 7.5)
    add_heading(doc, "Appendix B. Deliverable inventory", 2)
    for lead, text in [
        ("Raw data.", " CSV and JSON exports for projects, users, profiles, pair scores, supervisor scores, assignments, activities and runs."),
        ("Workbook.", " A consolidated 11-sheet XLSX with measured results and synthetic benchmark."),
        ("Charts.", " Fourteen publication-quality figures in 300-DPI PNG, SVG and PDF."),
        ("Reports.", " HTML, DOCX and PDF versions of this report."),
        ("Application.", " Thirteen administrator research pages and secured APIs."),
        ("Tests.", " Python, backend and database integrity checks."),
    ]:
        add_paragraph(doc, lead + text, bold_lead=lead)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def html_table(headers: list[str], rows: list[list[Any]]) -> str:
    head = "".join(f"<th>{html.escape(str(value))}</th>" for value in headers)
    body = "".join("<tr>" + "".join(f"<td>{html.escape(str(value))}</td>" for value in row) + "</tr>" for row in rows)
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def build_html(data: dict[str, Any], charts_dir: Path, output: Path) -> None:
    model_rows = [[row["label"], row["count"], f'{row["mean"]:.2f}', f'{row["median"]:.2f}', row["low"], row["medium"], row["high"]] for row in data["model_stats"]]
    assignment = data["assignment_stats"]
    images = [
        ("charts/figure-01-similarity-distribution.png", "Similarity score distributions"),
        ("charts/figure-04-model-correlation.png", "Pairwise model-score correlation"),
        ("charts/figure-08-supervisor-heatmap.png", "Project-to-supervisor match heatmap"),
        ("charts/figure-13-workload-after.png", "Workload after balancing"),
    ]
    figure_html = "".join(f'<figure><img src="{src}" alt="{html.escape(caption)}"><figcaption>{html.escape(caption)}</figcaption></figure>' for src, caption in images)
    sections = [
        ("Abstract", "All three model families completed full recorded-data coverage. Results are descriptive because the workbook contains no human ground truth."),
        ("Architecture and scope", "The experiment extends the existing React, Express, Flask and MongoDB application. Untagged application data were preserved."),
        ("Dataset and cleaning", "The Project Test Dataset worksheet produced 78 valid records; Notes was excluded; zero duplicates and zero invalid rows were found."),
        ("Ground truth and annotation", "The system supports independent 0/1/2 labels, continuous human scores and supervisor relevance grades. No human labels existed at report time."),
        ("Experimental design", "All 3,003 unique project pairs and all 1,560 project-supervisor combinations were measured for TF-IDF, multilingual Sentence-BERT and BGE-M3."),
        ("Similarity results", html_table(["Model", "Pairs", "Mean", "Median", "Low", "Medium", "High"], model_rows)),
        ("Supervisor assignment", f'All 78 projects received provisional assignments across {assignment["supervisors_used"]} supervisors; capacity violations: {assignment["capacity_exceeded"]}.'),
        ("Validity boundary", "Accuracy, precision, recall, F1, ROC-AUC, PR-AUC, MRR, MAP, NDCG, statistical error tests and a best-model decision are withheld until expert consensus labels exist."),
        ("Deployment recommendation", "Use TF-IDF as the explainable baseline, Sentence-BERT in semantic shadow mode, and BGE-M3 as an offline candidate until human validation justifies its resource cost."),
        ("Reproducibility", "Random seed 42; exact model names, dimensions, run identifiers, configuration, timings and outputs are retained in the research database and raw exports."),
    ]
    section_html = ""
    for title, content in sections:
        section_html += f"<section><h2>{html.escape(title)}</h2>{content if content.startswith('<table') else f'<p>{html.escape(content)}</p>'}</section>"
    document = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Comparative Evaluation of TF-IDF, Sentence-BERT and BGE-M3</title>
<style>
body{{font-family:Arial,sans-serif;color:#1f2937;max-width:1100px;margin:auto;padding:32px;line-height:1.55}}
header{{background:#17365d;color:white;padding:42px;border-radius:14px}} h1{{margin:0;font-size:34px}} h2{{color:#2e74b5;margin-top:34px}}
.warning{{background:#fff2cc;border-left:5px solid #bf9000;padding:16px;margin:24px 0}} table{{border-collapse:collapse;width:100%;font-size:14px}}
th{{background:#17365d;color:white}} th,td{{padding:9px;border:1px solid #d9e2f3;text-align:left}} figure{{margin:36px 0}} img{{max-width:100%;height:auto}} figcaption{{color:#667085;text-align:center;font-size:13px}}
footer{{margin-top:50px;border-top:1px solid #d9e2f3;padding-top:16px;color:#667085}}
</style></head><body>
<header><p>ACADEMIC RESEARCH REPORT</p><h1>Comparative Evaluation of TF-IDF, Sentence-BERT and BGE-M3</h1><p>Academic Project Similarity Detection and Supervisor Assignment</p><p>Hormuud University | Faculty of Computer Science and IT | 24 July 2026</p></header>
<div class="warning"><strong>Validity status:</strong> descriptive results complete; human ground-truth validation pending.</div>
{section_html}{figure_html}
<footer>Generated from measured research database exports. Synthetic benchmark labels are never presented as human validated.</footer>
</body></html>"""
    output.write_text(document, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--raw-dir", type=Path, required=True)
    parser.add_argument("--charts-dir", type=Path, required=True)
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--html", type=Path, required=True)
    args = parser.parse_args()
    data = build_data(args.raw_dir)
    build_docx(data, args.charts_dir, args.docx)
    build_html(data, args.charts_dir, args.html)
    print(json.dumps({"docx": str(args.docx), "html": str(args.html), "sections": 32}))


if __name__ == "__main__":
    main()
