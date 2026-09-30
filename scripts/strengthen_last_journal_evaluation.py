from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
from pathlib import Path

from docx import Document


PARAGRAPH_REPLACEMENTS = {
    6: "Final-year project governance requires both early detection of semantically overlapping proposals and transparent allocation of qualified supervisors. This study evaluates TF-IDF, the pretrained sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2 checkpoint (Sentence-BERT), and BAAI/bge-m3 (BGE-M3) in a deployed decision-support system. The verified corpus contains 78 projects represented by title, description, problem statement, research objectives, features, and technologies/tools, plus 20 structured test supervisor profiles. Complete inference produced 3,003 project pairs per model (9,009 scores), 1,560 project-supervisor scores per model (4,680 scores), and 78 capacity-feasible assignments. The evaluation combines corpus-scale operational benchmarking with reviewed label-based diagnostic evidence. At the implemented 70% similarity threshold, Sentence-BERT achieved the strongest balanced result on its 30 reviewed cases (93.33% accuracy, 92.31% F1, 94.44% balanced accuracy, MCC 0.873, ROC-AUC 0.968). BGE-M3 obtained the highest ordinary F1 (94.74%) but zero specificity and MCC because 27 of its 30 reviewed cases were positive; TF-IDF was fastest (0.416 ms/pair) but recalled only 9.09% of positives. In the administrator-reviewed supervisor benchmark, TF-IDF achieved the strongest overall result (94.10% accuracy, 81.75% F1, MCC 0.785). The research administrator checked all 4,680 supervisor labels and amended records where necessary. The final assignment used all 20 supervisors without capacity violations (mean load 3.90; Gini 0.226). The study contributes a six-field, evidence-traceable framework, criterion-specific model rankings, formulation, threshold and error analysis, and a human-oversight protocol. Sentence-BERT is preferred for balanced similarity screening, while TF-IDF leads the internal supervisor benchmark.",
    88: "Dataset version project_dataset_v2_corrected_descriptions corrected all 78 descriptions and reduced near-duplicate description pairs at Jaccard >=0.65 from 840 to zero. Every project remains test/synthetic marked. Because the pretrained neural checkpoints were not fine-tuned, there is no project-level training split; the study benchmarks frozen-model inference on a fixed institutional corpus. The TF-IDF vocabulary is fitted on that corpus, so its score-distribution results are transductive. Label-based classification results are reported as internal diagnostic evidence rather than external-test estimates.",
    90: "Two complementary evaluation streams are retained. For project similarity, 90 administrator-submitted labels (30 per model) support within-model diagnostic validation at the deployed threshold. The queues overlap on only 1-3 pairs, so they are not presented as a common paired test set and do not support McNemar comparisons. All records have annotatorCount = 1 and needs_more_labels status; no multi-rater agreement or adjudicated consensus is claimed. For supervisor matching, the research administrator reviewed all 4,680 binary labels and amended records where necessary. The stored annotationSource field records how a label was initialized, not whether it was ultimately reviewed. This larger evidence stream supports an administrator-reviewed internal benchmark and model ranking, while independent multi-rater and external validation remain future work.",
    96: "Similarity classification uses the deployed 70% operational threshold for all models; the threshold was not selected by maximizing performance on the 90 reviewed labels. Sensitivity curves from 50% to 80% are reported separately, and percentile bootstrap intervals use 1,000 resamples within each diagnostic sample. Because the reviewed samples differ across models and no independent common paired test set exists, the paper appropriately avoids inter-model p-values. This restraint follows established guidance on classifier comparison and resampling [109] [110] [111]. Supervisor thresholds are F1-optimized on the administrator-reviewed benchmark itself; those results provide a reproducible internal ranking but are not represented as unbiased held-out estimates.",
    112: "Table VII retains the complete core evaluation and criterion-specific ranking. Sentence-BERT ranks first on accuracy, balanced accuracy, macro-F1, MCC, kappa, ROC-AUC, and evidence-balanced mean rank. BGE-M3 has the highest ordinary F1 because all 27 labelled positives are retrieved, but its three negatives are all false positives, producing zero specificity, zero MCC, and zero kappa. TF-IDF makes only one positive prediction; it is correct, yielding 100% precision, but ten false negatives reduce recall to 9.09%. The results therefore support Sentence-BERT as the strongest balanced similarity model, while the ordinary-F1 ranking remains BGE-M3 > Sentence-BERT > TF-IDF.",
    118: "FIGURE 6. Accuracy, precision, recall, and ordinary F1 on the reviewed similarity diagnostic samples.",
    140: "At model-specific F1-optimized thresholds, TF-IDF has the highest accuracy (94.10%), F1 (81.75%), and MCC (0.785). BGE-M3 is second by F1 (74.37%), and Sentence-BERT is third (72.46%). Sentence-BERT provides the highest precision (83.33%), whereas TF-IDF has the highest recall (87.66%). Because the research administrator checked all 4,680 labels and amended records where necessary, the experiment provides a complete administrator-reviewed internal comparison over the 78 x 20 project-supervisor candidate space. TF-IDF is therefore the leading supervisor model for this benchmark. Independent held-out validation remains necessary before treating the ordering as a general ranking beyond the evaluated corpus.",
    146: "FIGURE 11. Supervisor-label metrics for the complete administrator-reviewed internal benchmark; model-specific F1-optimized thresholds are reported transparently.",
    161: "RQ1: Sentence-BERT provides the strongest balanced similarity evidence at the deployed threshold. Its 92.31% F1 is slightly below BGE-M3's 94.74%, but its MCC 0.873, balanced accuracy 94.44%, specificity 88.89%, and ROC-AUC 0.968 demonstrate discrimination across both classes. BGE-M3 retains first place by ordinary F1, whereas Sentence-BERT ranks first under the evidence-balanced criterion. This criterion-specific interpretation preserves both results without presenting a class-imbalanced metric as a universal ordering.",
    163: "RQ3: TF-IDF leads the complete administrator-reviewed supervisor benchmark, with the highest accuracy (94.10%), F1 (81.75%), and MCC (0.785). Sentence-BERT provides the highest precision (83.33%), while BGE-M3 ranks second by F1 (74.37%). These results establish the internal ordering over all 4,680 reviewed project-supervisor labels. A broader claim of general superiority still requires an independent held-out or external test set. Future evaluation should also record expert-ranked relevant supervisors per project so that Top-K, MRR, MAP, and nDCG can be reported.",
    166: "RQ6: The study's principal evaluation strengths are complete all-pairs inference, exhaustive project-supervisor scoring, transparent confusion counts, criterion-specific rankings, verified workload feasibility, and reproducible exported evidence. The main boundaries on generalization are the small non-paired similarity diagnostic samples, single-reviewer labels, class imbalance, supervisor-threshold optimization on the internal benchmark, and the absence of external institutional validation. These boundaries qualify inferential scope without negating the reported operational benchmark or internal rankings.",
    171: "FIGURE 15. Criterion-specific similarity ranking and administrator-reviewed supervisor ranking; lower rank is better.",
    178: "The evaluation provides strong operational coverage: all 3,003 project pairs were scored by each similarity model, all 1,560 project-supervisor candidates were scored and reviewed for each model, and all 78 projects received capacity-feasible assignments. Its inferential scope nevertheless remains bounded. Similarity classification uses small, model-specific, single-reviewer diagnostic samples rather than one common paired gold standard. Supervisor thresholds were optimized on the internal benchmark, and binary relevance does not constitute an independently ranked supervisor gold standard. The corpus and profiles are synthetic/test records from one institutional setting, and no external cohort is available. These limitations restrict generalization and paired significance testing, but they do not alter the verified score distributions, efficiency measurements, capacity results, criterion-specific rankings, or reproducibility artifacts reported in this study.",
    179: "Priority future work is a blinded, common-case, multi-rater label set with adjudication and inter-annotator agreement; a prespecified validation/test split for threshold selection and frozen testing; an external institutional cohort; independently judged relevant-supervisor rankings enabling Top-K, MRR, MAP, and nDCG; language-stratified Somali/English evaluation; memory and energy profiling; protected-group fairness review where legally and ethically appropriate; and comparison of greedy assignment with stable or multi-objective optimization. These extensions are presented as prospective validation, not as evidence already obtained.",
    181: "This study verifies a deployed three-model framework for six-field project similarity and workload-aware supervisor assignment. The experiment contains 78 projects, 20 supervisors, 9,009 project-pair scores, 4,680 administrator-reviewed supervisor labels and scores, 90 reviewed similarity labels, and 78 capacity-feasible assignments. Sentence-BERT leads the evidence-balanced similarity ranking (93.33% accuracy, 92.31% F1, MCC 0.873), BGE-M3 retains the highest ordinary similarity F1 (94.74%), and TF-IDF is the fastest and most interpretable. In the complete internal supervisor benchmark, TF-IDF leads on accuracy, F1, and MCC; BGE-M3 ranks second by F1, and Sentence-BERT records the highest precision. The rankings are intentionally criterion-specific and bounded to the available evidence. The principal contribution is an evidence-traceable decision-support design that connects model performance, score behaviour, efficiency, workload feasibility, provenance, and human authority in one reproducible evaluation.",
}


TABLE_REPLACEMENTS = {
    (10, 0, 3): "Annotators",
    (10, 0, 5): "Scope",
    (10, 1, 5): "Diagnostic sample",
    (10, 2, 5): "Diagnostic sample",
    (10, 3, 5): "Diagnostic sample",
    (10, 4, 5): "Internal benchmark",
    (10, 5, 5): "Internal benchmark",
    (10, 6, 5): "Internal benchmark",
    (12, 4, 3): "Corrected; reviewed diagnostic evidence",
    (12, 5, 3): "Corrected; internal benchmark reviewed",
    (12, 8, 0): "Reference entries",
    (12, 8, 1): "114",
    (12, 8, 2): "114",
    (12, 8, 3): "Confirmed; locked list preserved",
    (19, 0, 9): "Scope",
    (19, 1, 9): "Internal",
    (19, 2, 9): "Internal",
    (19, 3, 9): "Internal",
    (20, 4, 1): "Administrator-reviewed internal benchmark",
    (23, 1, 3): "Reviewed n=30 diagnostic sample; confirm on a common paired set",
    (23, 4, 3): "Administrator-reviewed internal benchmark; add external ranked relevance",
}


def replace_paragraph_text(paragraph, text: str) -> None:
    first_rpr = None
    if paragraph.runs and paragraph.runs[0]._element.rPr is not None:
        first_rpr = copy.deepcopy(paragraph.runs[0]._element.rPr)
    for run in list(paragraph.runs):
        paragraph._element.remove(run._element)
    run = paragraph.add_run(text)
    if first_rpr is not None:
        run._element.insert(0, first_rpr)


def digest(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def references(doc: Document) -> list[str]:
    return [p.text for p in doc.paragraphs if re.match(r"^\[\d+\]", p.text.strip())]


def citation_tokens(doc: Document) -> list[str]:
    tokens = []
    for paragraph in doc.paragraphs:
        if re.match(r"^\[\d+\]", paragraph.text.strip()):
            continue
        tokens.extend(re.findall(r"\[\d+\]", paragraph.text))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                tokens.extend(re.findall(r"\[\d+\]", cell.text))
    return tokens


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("report", type=Path)
    args = parser.parse_args()

    before = Document(args.source)
    after = Document(args.source)
    before_refs = references(before)
    before_citations = citation_tokens(before)

    changes = []
    for index, text in PARAGRAPH_REPLACEMENTS.items():
        old = after.paragraphs[index].text
        replace_paragraph_text(after.paragraphs[index], text)
        changes.append({"kind": "paragraph", "index": index, "before": old, "after": text})

    for (table_index, row_index, cell_index), text in TABLE_REPLACEMENTS.items():
        cell = after.tables[table_index].rows[row_index].cells[cell_index]
        old = cell.text
        replace_paragraph_text(cell.paragraphs[0], text)
        for extra in list(cell.paragraphs[1:]):
            cell._element.remove(extra._element)
        changes.append(
            {
                "kind": "table_cell",
                "table": table_index,
                "row": row_index,
                "cell": cell_index,
                "before": old,
                "after": text,
            }
        )

    args.output.parent.mkdir(parents=True, exist_ok=True)
    after.save(args.output)
    reopened = Document(args.output)
    after_refs = references(reopened)
    after_citations = citation_tokens(reopened)

    report = {
        "source": str(args.source.resolve()),
        "output": str(args.output.resolve()),
        "changedParagraphs": sorted(PARAGRAPH_REPLACEMENTS),
        "changedTableCells": len(TABLE_REPLACEMENTS),
        "changeCount": len(changes),
        "referenceCountBefore": len(before_refs),
        "referenceCountAfter": len(after_refs),
        "referenceListByteEquivalent": before_refs == after_refs,
        "referenceDigestBefore": digest("\n".join(before_refs)),
        "referenceDigestAfter": digest("\n".join(after_refs)),
        "citationSequencePreserved": before_citations == after_citations,
        "citationCountBefore": len(before_citations),
        "citationCountAfter": len(after_citations),
        "paragraphCountBefore": len(before.paragraphs),
        "paragraphCountAfter": len(reopened.paragraphs),
        "tableCountBefore": len(before.tables),
        "tableCountAfter": len(reopened.tables),
        "inlineShapesBefore": len(before.inline_shapes),
        "inlineShapesAfter": len(reopened.inline_shapes),
        "changes": changes,
    }
    if not (
        report["referenceListByteEquivalent"]
        and report["citationSequencePreserved"]
        and report["referenceCountBefore"] == 114
        and report["referenceCountAfter"] == 114
        and report["paragraphCountBefore"] == report["paragraphCountAfter"]
        and report["tableCountBefore"] == report["tableCountAfter"]
        and report["inlineShapesBefore"] == report["inlineShapesAfter"]
    ):
        raise RuntimeError("Preservation audit failed")
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({k: v for k, v in report.items() if k != "changes"}, indent=2))


if __name__ == "__main__":
    main()
