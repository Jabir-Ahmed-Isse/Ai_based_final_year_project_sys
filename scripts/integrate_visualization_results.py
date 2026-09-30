from __future__ import annotations

import csv
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from PIL import Image

from build_research_report import MUTED, add_paragraph, set_font


ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "outputs" / "journal-package-20260724"
SOURCE = PACKAGE / "final-journal-manuscript.docx"
OUTPUT = PACKAGE / "final-journal-manuscript-complete-visuals.docx"
VISUALS = PACKAGE / "system-visualizations"
INVENTORY = VISUALS / "visualization-capture-inventory.csv"
REPORT = PACKAGE / "visual-integration-report.json"


def load_inventory() -> list[dict[str, str]]:
    with INVENTORY.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def fitted_width(path: Path, maximum_width: float = 6.25, maximum_height: float = 7.0) -> float:
    with Image.open(path) as image:
        width_px, height_px = image.size
    width = min(maximum_width, maximum_height * width_px / height_px)
    if width_px < 400:
        width = min(width, 3.2)
    elif width_px < 700:
        width = min(width, 5.3)
    return max(1.4, width)


def move_before(paragraph, target) -> None:
    target._p.addprevious(paragraph._p)


def insert_heading(document: Document, target, text: str, level: int) -> None:
    paragraph = document.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    move_before(paragraph, target)


def insert_text(document: Document, target, text: str) -> None:
    paragraph = add_paragraph(document, text)
    move_before(paragraph, target)


def insert_figure(document: Document, target, image_path: Path, caption: str) -> None:
    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.keep_with_next = True
    run = image_paragraph.add_run()
    shape = run.add_picture(str(image_path), width=Inches(fitted_width(image_path)))
    shape._inline.docPr.set("descr", caption)
    move_before(image_paragraph, target)

    caption_paragraph = document.add_paragraph(style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.keep_with_next = False
    caption_run = caption_paragraph.add_run(caption)
    set_font(caption_run, size=9, color=MUTED, italic=True)
    move_before(caption_paragraph, target)


def row_map(rows: list[dict[str, str]]) -> dict[int, list[dict[str, str]]]:
    grouped: dict[int, list[dict[str, str]]] = {}
    for row in rows:
        grouped.setdefault(int(row["visual_number"]), []).append(row)
    for visual_rows in grouped.values():
        visual_rows.sort(key=lambda row: int(row["part"]))
    return grouped


def main() -> None:
    rows = load_inventory()
    grouped = row_map(rows)
    if sorted(grouped) != list(range(1, 18)):
        raise RuntimeError(f"Expected visual sections 1-17, found {sorted(grouped)}")

    document = Document(SOURCE)
    target_matches = [
        paragraph for paragraph in document.paragraphs if paragraph.text.strip() == "5. Discussion"
    ]
    if len(target_matches) != 1:
        raise RuntimeError(f"Expected one '5. Discussion' heading, found {len(target_matches)}")
    target = target_matches[0]

    insert_heading(document, target, "4.8 Live visualization-page analysis", 2)
    insert_text(
        document,
        target,
        "The administrator Visualization page was captured directly from the running system "
        "after the recorded experiment was loaded. The panels below are placed inside the "
        "Results section, rather than after the references, because they visualize the same "
        "78-project evidence analysed in Sections 4.1-4.7. They remain descriptive: higher "
        "scores or stronger agreement do not establish predictive accuracy without expert labels.",
    )

    figure_number = 14
    overview = VISUALS / "visualization-page-overview.png"
    insert_figure(
        document,
        target,
        overview,
        f"Figure {figure_number}. Live administrator Research Visualizations page overview for "
        "the recorded 78-project experiment.",
    )
    figure_number += 1

    groups = [
        (
            "4.8.1 Experiment coverage and completion",
            "The live summary confirms 78 recorded projects, 9,009 project-model score records "
            "(3,003 canonical project pairs per completed model), 4,680 supervisor-match records "
            "(1,560 candidates per completed model), 78 provisional assignments and zero "
            "expert-consensus annotations.",
            [1, 2, 3, 4, 5],
        ),
        (
            "4.8.2 Project-similarity distributions and descriptive statistics",
            "The distribution, risk-band and confidence-interval panels expose model calibration "
            "differences across the same persisted pair set. These visual differences are useful "
            "for threshold inspection but are not accuracy rankings.",
            [6, 7, 11],
        ),
        (
            "4.8.3 Field contribution, agreement and disagreement",
            "The field panel retains all six faculty-approved project fields. The correlation and "
            "largest-disagreement panels show where the three scoring systems converge or diverge "
            "on identical project pairs.",
            [8, 10, 12],
        ),
        (
            "4.8.4 Measured processing evidence",
            "The latency and processing-summary panels reproduce the persisted timing metadata. "
            "The values are implementation- and cache-context measurements and therefore support "
            "operational reporting rather than a hardware-neutral speed benchmark.",
            [9, 15],
        ),
        (
            "4.8.5 Supervisor matching, workload and capacity",
            "The supervisor panels summarize semantic and workload-adjusted candidate scores and "
            "the final capacity-aware allocation across all 20 synthetic supervisor profiles.",
            [13, 14],
        ),
        (
            "4.8.6 Human-validation boundary",
            "The live validation panel correctly reports every label-dependent measurement as "
            "unavailable because the expert-consensus sample is zero. Retaining this status inside "
            "the Results section prevents descriptive model scores from being misreported as "
            "accuracy, F1, AUC or statistical superiority.",
            [16, 17],
        ),
    ]

    for heading, analysis, visual_numbers in groups:
        insert_heading(document, target, heading, 3)
        insert_text(document, target, analysis)
        for visual_number in visual_numbers:
            for row in grouped[visual_number]:
                image_path = VISUALS / row["cropped_file"]
                part = int(row["part"])
                parts = int(row["parts"])
                part_text = "" if parts == 1 else f" Part {part} of {parts}."
                caption = (
                    f"Figure {figure_number}. Live Visualization page analysis: "
                    f"{row['live_page_title']}.{part_text}"
                )
                insert_figure(document, target, image_path, caption)
                figure_number += 1

    if figure_number != 34:
        raise RuntimeError(f"Expected to finish at Figure 33, finished at Figure {figure_number - 1}")

    document.save(OUTPUT)

    reopened = Document(OUTPUT)
    paragraphs = [paragraph.text.strip() for paragraph in reopened.paragraphs]
    result_index = paragraphs.index("4.8 Live visualization-page analysis")
    discussion_index = paragraphs.index("5. Discussion")
    references_index = paragraphs.index("References")

    with zipfile.ZipFile(OUTPUT) as archive:
        names = set(archive.namelist())
        media = [name for name in names if name.startswith("word/media/")]
        xml = archive.read("word/document.xml").decode("utf-8")
        alt_text_count = len(
            re.findall(r"<wp:docPr\b[^>]*\bdescr=\"[^\"]+\"", xml)
        )

    report = {
        "output": str(OUTPUT.resolve()),
        "source_preserved": str(SOURCE.resolve()),
        "live_visual_sections": len(grouped),
        "live_visual_images": len(rows),
        "live_overview_images": 1,
        "new_result_figures": 20,
        "first_new_figure": 14,
        "last_new_figure": 33,
        "embedded_media_total": len(media),
        "drawing_alt_text_total": alt_text_count,
        "inserted_inside_results": result_index < discussion_index < references_index,
        "system_interface_appendix_present": any(
            text.startswith("Appendix C. Complete live system-interface") for text in paragraphs
        ),
        "visualization_appendix_present": any(
            text.startswith("Appendix D. Complete live research-visualization") for text in paragraphs
        ),
        "comparative_evaluation_preserved": (
            "4.4 Comparative evaluation and evidence-bounded rankings" in paragraphs
        ),
    }
    REPORT.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
