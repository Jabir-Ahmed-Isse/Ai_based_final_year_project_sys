"""Read-only audit for the symbolic-mathematics journal edition."""

from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "outputs" / "journal-package-20260724"
MANUSCRIPT = PACKAGE / "final-journal-manuscript-symbolic-mathematics.docx"
REPORT = PACKAGE / "symbolic-mathematics-final-audit.json"


document = Document(MANUSCRIPT)
paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
styles = [paragraph.style.name for paragraph in document.paragraphs]

main_figure_captions = [text for text in paragraphs if re.match(r"Figure \d+\.", text)]
table_captions = [text for text in paragraphs if re.match(r"Table \d+\.", text)]

reference_start = paragraphs.index("References") + 1
appendix_start = paragraphs.index("Appendix A. Journal-quality pseudocode")
reference_numbers = [
    int(match.group(1))
    for text in paragraphs[reference_start:appendix_start]
    if (match := re.match(r"\[(\d+)\]\s", text))
]

with zipfile.ZipFile(MANUSCRIPT) as archive:
    names = set(archive.namelist())
    xml = archive.read("word/document.xml").decode("utf-8")
    media = [name for name in names if name.startswith("word/media/")]
    alt_text_count = len(
        re.findall(r"<wp:docPr\b[^>]*\bdescr=\"[^\"]+\"", xml)
    )
    equation_objects = len(re.findall(r"<m:oMath\b", xml))
    equation_paragraphs = len(re.findall(r"<m:oMathPara\b", xml))
    fractions = len(re.findall(r"<m:f\b", xml))
    radicals = len(re.findall(r"<m:rad\b", xml))
    subscripts = len(re.findall(r"<m:sSub\b", xml))
    superscripts = len(re.findall(r"<m:sSup\b", xml))
    matrices = len(re.findall(r"<m:m\b", xml))
    package_integrity = {
        "[Content_Types].xml",
        "word/document.xml",
        "word/styles.xml",
        "word/settings.xml",
    }.issubset(names)

suspicious_sequences = {
    sequence: sum(text.count(sequence) for text in paragraphs)
    for sequence in ("\ufffd", "Ãƒ", "Ã¢â‚¬", "ÃŽ")
}

math_headings = [
    f"3.11.{number} {title}"
    for number, title in [
        (1, "Corpus size and experimental cardinality"),
        (2, "TF-IDF and vector similarity"),
        (3, "Transformer representation and multi-field aggregation"),
        (4, "Risk bands and workload-aware supervisor ranking"),
        (5, "Descriptive and agreement statistics"),
        (6, "Human-label-dependent evaluation metrics"),
    ]
]

plain_formula_prefixes = (
    "Term frequency: tf(",
    "Inverse document frequency: idf(",
    "TF-IDF weight: w(",
    "Cosine similarity: cos(",
    "Weighted multi-field similarity: S =",
    "Risk: Low if S<",
)

checks = {
    "manuscript_reopened": True,
    "open_xml_integrity": package_integrity,
    "table_count": len(document.tables),
    "table_caption_count": len(table_captions),
    "main_figure_caption_count": len(main_figure_captions),
    "embedded_media_count": len(media),
    "drawing_alt_text_count": alt_text_count,
    "omml_equation_objects": equation_objects,
    "omml_equation_paragraphs": equation_paragraphs,
    "omml_fraction_count": fractions,
    "omml_radical_count": radicals,
    "omml_subscript_count": subscripts,
    "omml_superscript_count": superscripts,
    "omml_matrix_count": matrices,
    "all_math_subheadings_present": all(heading in paragraphs for heading in math_headings),
    "plain_formula_list_removed": not any(
        text.startswith(prefix)
        for text in paragraphs
        for prefix in plain_formula_prefixes
    ),
    "reference_number_sequence_valid": reference_numbers == list(range(1, 161)),
    "comparative_evaluation_preserved": (
        "4.4 Comparative evaluation and evidence-bounded rankings" in paragraphs
    ),
    "visualization_results_preserved": (
        "4.8 Live visualization-page analysis" in paragraphs
    ),
    "mathematics_precedes_architecture": (
        paragraphs.index("3.11 Mathematical formulation")
        < paragraphs.index("3.12 Implementation architecture")
    ),
    "gated_metric_warning_present": any(
        "With zero expert-consensus labels" in text for text in paragraphs
    ),
    "suspicious_text_sequences": suspicious_sequences,
}

checks["passed"] = all(
    [
        checks["open_xml_integrity"],
        checks["table_count"] == 11,
        checks["table_caption_count"] == 11,
        checks["main_figure_caption_count"] == 33,
        checks["embedded_media_count"] == 34,
        checks["drawing_alt_text_count"] == 34,
        checks["omml_equation_objects"] == 30,
        checks["omml_equation_paragraphs"] == 30,
        checks["omml_fraction_count"] >= 20,
        checks["omml_radical_count"] >= 4,
        checks["omml_subscript_count"] >= 50,
        checks["omml_superscript_count"] >= 8,
        checks["omml_matrix_count"] >= 1,
        checks["all_math_subheadings_present"],
        checks["plain_formula_list_removed"],
        checks["reference_number_sequence_valid"],
        checks["comparative_evaluation_preserved"],
        checks["visualization_results_preserved"],
        checks["mathematics_precedes_architecture"],
        checks["gated_metric_warning_present"],
        not any(suspicious_sequences.values()),
    ]
)

REPORT.write_text(json.dumps(checks, indent=2, ensure_ascii=False), encoding="utf-8")
print(json.dumps(checks, indent=2, ensure_ascii=True))
