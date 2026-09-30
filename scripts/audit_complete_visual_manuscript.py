"""Read-only audit of the complete-visuals journal manuscript."""

from __future__ import annotations

import csv
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from PIL import Image, ImageStat


ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "outputs" / "journal-package-20260724"
MANUSCRIPT = PACKAGE / "final-journal-manuscript-complete-visuals.docx"
LIVE_INVENTORY = PACKAGE / "system-visualizations" / "visualization-capture-inventory.csv"
REPORT = PACKAGE / "complete-visuals-audit.json"


def load_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def image_check(path: Path) -> dict[str, object]:
    with Image.open(path) as image:
        grayscale = image.convert("L")
        deviation = float(ImageStat.Stat(grayscale).stddev[0])
        return {
            "file": path.name,
            "width": image.width,
            "height": image.height,
            "grayscale_standard_deviation": round(deviation, 2),
            "nonblank": image.width >= 200 and image.height >= 100 and deviation >= 3,
        }


document = Document(MANUSCRIPT)
paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
main_figure_captions = [text for text in paragraphs if re.match(r"Figure \d+\.", text)]
table_captions = [text for text in paragraphs if re.match(r"Table \d+\.", text)]
new_result_captions = [
    text
    for text in paragraphs
    if re.match(r"Figure (1[4-9]|2[0-9]|3[0-3])\.", text)
]

reference_start = paragraphs.index("References") + 1
appendix_start = paragraphs.index("Appendix A. Journal-quality pseudocode")
reference_numbers = [
    int(match.group(1))
    for text in paragraphs[reference_start:appendix_start]
    if (match := re.match(r"\[(\d+)\]\s", text))
]

live_rows = load_csv(LIVE_INVENTORY)
live_images = [
    image_check(PACKAGE / "system-visualizations" / row["cropped_file"])
    for row in live_rows
]
overview_check = image_check(
    PACKAGE / "system-visualizations" / "visualization-page-overview.png"
)

with zipfile.ZipFile(MANUSCRIPT) as archive:
    names = set(archive.namelist())
    media = sorted(name for name in names if name.startswith("word/media/"))
    document_xml = archive.read("word/document.xml").decode("utf-8")
    alt_text_count = len(
        re.findall(r"<wp:docPr\b[^>]*\bdescr=\"[^\"]+\"", document_xml)
    )
    package_integrity = {
        "[Content_Types].xml",
        "word/document.xml",
        "word/styles.xml",
        "word/settings.xml",
    }.issubset(names)

suspicious_sequences = {
    sequence: sum(text.count(sequence) for text in paragraphs)
    for sequence in ("\ufffd", "Ãƒ", "Ã¢â‚¬", "ÃŽ")
}

checks = {
    "manuscript_reopened": True,
    "open_xml_integrity": package_integrity,
    "table_count": len(document.tables),
    "table_caption_count": len(table_captions),
    "main_figure_caption_count": len(main_figure_captions),
    "new_result_figure_caption_count": len(new_result_captions),
    "embedded_media_count": len(media),
    "drawing_alt_text_count": alt_text_count,
    "live_visual_section_count": len({int(row["visual_number"]) for row in live_rows}),
    "live_visual_image_count": len(live_rows),
    "live_overview_nonblank": overview_check["nonblank"],
    "live_images_nonblank": all(item["nonblank"] for item in live_images),
    "reference_number_sequence_valid": reference_numbers == list(range(1, 161)),
    "comparative_evaluation_preserved": (
        "4.4 Comparative evaluation and evidence-bounded rankings" in paragraphs
    ),
    "results_visualization_section_present": (
        "4.8 Live visualization-page analysis" in paragraphs
    ),
    "results_visualization_section_precedes_discussion_and_references": (
        paragraphs.index("4.8 Live visualization-page analysis")
        < paragraphs.index("5. Discussion")
        < paragraphs.index("References")
    ),
    "system_interface_appendix_absent": (
        "Appendix C. Complete live system-interface evidence" not in paragraphs
    ),
    "visualization_appendix_absent": (
        "Appendix D. Complete live research-visualization evidence" not in paragraphs
    ),
    "suspicious_text_sequences": suspicious_sequences,
    "live_images": live_images,
    "live_overview": overview_check,
}

checks["passed"] = all(
    [
        checks["open_xml_integrity"],
        checks["table_count"] == 11,
        checks["table_caption_count"] == 11,
        checks["main_figure_caption_count"] == 33,
        checks["new_result_figure_caption_count"] == 20,
        checks["embedded_media_count"] == 34,
        checks["drawing_alt_text_count"] == 34,
        checks["live_visual_section_count"] == 17,
        checks["live_visual_image_count"] == 19,
        checks["live_overview_nonblank"],
        checks["live_images_nonblank"],
        checks["reference_number_sequence_valid"],
        checks["comparative_evaluation_preserved"],
        checks["results_visualization_section_present"],
        checks["results_visualization_section_precedes_discussion_and_references"],
        checks["system_interface_appendix_absent"],
        checks["visualization_appendix_absent"],
        not any(suspicious_sequences.values()),
    ]
)

REPORT.write_text(json.dumps(checks, indent=2, ensure_ascii=False), encoding="utf-8")
print(json.dumps({key: value for key, value in checks.items() if not key.endswith("images")}, indent=2))
