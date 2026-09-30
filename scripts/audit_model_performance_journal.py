from __future__ import annotations

import argparse
import json
import re
import zipfile
from pathlib import Path

from docx import Document


def all_text(document: Document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def audit_docx(path: Path) -> dict:
    document = Document(path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
    text = all_text(document)
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        xml = archive.read("word/document.xml").decode("utf-8")
        media = [name for name in names if name.startswith("word/media/")]
        package_integrity = {
            "[Content_Types].xml",
            "word/document.xml",
            "word/styles.xml",
            "word/settings.xml",
        }.issubset(names)
    return {
        "path": str(path),
        "reopened": True,
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "inline_shapes": len(document.inline_shapes),
        "media_files": len(media),
        "package_integrity": package_integrity,
        "omml_equation_objects": len(re.findall(r"<m:oMath\b", xml)),
        "omml_equation_paragraphs": len(re.findall(r"<m:oMathPara\b", xml)),
        "table_captions": [
            paragraph for paragraph in paragraphs if re.match(r"^Table \d+\.", paragraph)
        ],
        "figure_captions": [
            paragraph for paragraph in paragraphs if re.match(r"^Figure \d+\.", paragraph)
        ],
        "text": text,
        "paragraph_text": paragraphs,
        "document": document,
    }


def performance_table(document: Document):
    matches = [
        table
        for table in document.tables
        if [cell.text.strip() for cell in table.rows[0].cells]
        == ["Metric", "TF-IDF", "Sentence-BERT", "BGE-M3", "Interpretation"]
    ]
    if len(matches) != 1:
        raise ValueError(f"Expected one performance table; found {len(matches)}")
    return matches[0]


def metric_rows(table) -> dict[str, list[str]]:
    return {
        row.cells[0].text.strip(): [cell.text.strip() for cell in row.cells[1:]]
        for row in table.rows[1:]
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--journal", required=True, type=Path)
    parser.add_argument("--report", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    journal = audit_docx(args.journal)
    report = audit_docx(args.report)
    journal_table = performance_table(journal["document"])
    report_table = performance_table(report["document"])
    journal_metrics = metric_rows(journal_table)
    report_metrics = metric_rows(report_table)

    required_ne_metrics = ["Accuracy", "Precision", "Recall", "Macro F1", "MCC"]
    journal_ne = {
        metric: journal_metrics.get(metric, [])[:3] for metric in required_ne_metrics
    }
    report_ne = {
        metric: report_metrics.get(metric, [])[:3] for metric in required_ne_metrics
    }

    journal_caption_numbers = [
        int(re.match(r"^Table (\d+)\.", caption).group(1))
        for caption in journal["table_captions"]
    ]
    reference_start = journal["paragraph_text"].index("References") + 1
    appendix_start = journal["paragraph_text"].index(
        "Appendix A. Journal-quality pseudocode"
    )
    reference_numbers = [
        int(match.group(1))
        for paragraph in journal["paragraph_text"][reference_start:appendix_start]
        if (match := re.match(r"^\[(\d+)\]\s", paragraph))
    ]
    journal_text = journal["text"]
    report_text = report["text"]
    performance_heading_index = journal["paragraph_text"].index(
        "4.4.4 Metric-complete model-performance results"
    )
    efficiency_heading_index = journal["paragraph_text"].index("4.5 Efficiency")

    checks = {
        "journal_reopened": journal["reopened"],
        "report_reopened": report["reopened"],
        "journal_package_integrity": journal["package_integrity"],
        "report_package_integrity": report["package_integrity"],
        "journal_tables": journal["tables"],
        "report_tables": report["tables"],
        "journal_inline_shapes": journal["inline_shapes"],
        "journal_media_files": journal["media_files"],
        "journal_omml_objects": journal["omml_equation_objects"],
        "journal_omml_paragraphs": journal["omml_equation_paragraphs"],
        "journal_table_caption_numbers": journal_caption_numbers,
        "journal_table_captions_sequential": journal_caption_numbers
        == list(range(1, 13)),
        "journal_main_figure_captions": len(journal["figure_captions"]),
        "journal_references": len(reference_numbers),
        "journal_reference_sequence_valid": reference_numbers == list(range(1, 161)),
        "performance_section_before_efficiency": performance_heading_index
        < efficiency_heading_index,
        "journal_ne_metrics": journal_ne,
        "report_ne_metrics": report_ne,
        "journal_all_ne_metrics_complete": all(
            values == ["NE", "NE", "NE"] for values in journal_ne.values()
        ),
        "report_all_ne_metrics_complete": all(
            values == ["NE", "NE", "NE"] for values in report_ne.values()
        ),
        "journal_historical_metrics_present": all(
            value in journal_text for value in ["96.42%", "95.65%", "99.00%", "97.29%"]
        ),
        "report_historical_metrics_present": all(
            value in report_text for value in ["96.42%", "95.65%", "99.00%", "97.29%"]
        ),
        "journal_zero_label_statement_present": (
            "zero expert-labelled pairs" in journal_text
            and "Independent expert-labelled pairs" in journal_text
        ),
        "journal_predictive_quality_rank_withheld": (
            journal_metrics.get("Predictive-quality rank", [])[:3]
            == ["NE", "NE", "NE"]
        ),
        "journal_visualization_section_preserved": (
            "4.8 Live visualization-page analysis" in journal_text
        ),
        "journal_comparative_section_preserved": (
            "4.4 Comparative evaluation and evidence-bounded rankings" in journal_text
        ),
        "journal_references_preserved": "References" in journal_text,
        "suspicious_mojibake_count": sum(
            journal_text.count(sequence)
            for sequence in ("\ufffd", "ÃƒÆ’", "ÃƒÂ¢Ã¢â€šÂ¬", "ÃƒÅ½")
        ),
    }

    expected = {
        "journal_tables": 12,
        "report_tables": 13,
        "journal_inline_shapes": 34,
        "journal_media_files": 34,
        "journal_omml_objects": 30,
        "journal_omml_paragraphs": 30,
        "journal_main_figure_captions": 33,
        "journal_references": 160,
        "suspicious_mojibake_count": 0,
    }
    failures = []
    for key, value in expected.items():
        if checks.get(key) != value:
            failures.append(f"{key}: expected {value}, found {checks.get(key)}")
    for key in [
        "journal_package_integrity",
        "report_package_integrity",
        "journal_table_captions_sequential",
        "journal_reference_sequence_valid",
        "performance_section_before_efficiency",
        "journal_all_ne_metrics_complete",
        "report_all_ne_metrics_complete",
        "journal_historical_metrics_present",
        "report_historical_metrics_present",
        "journal_zero_label_statement_present",
        "journal_predictive_quality_rank_withheld",
        "journal_visualization_section_preserved",
        "journal_comparative_section_preserved",
        "journal_references_preserved",
    ]:
        if not checks.get(key):
            failures.append(f"{key}: failed")

    output = {
        "status": "passed" if not failures else "failed",
        "checks": checks,
        "failures": failures,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(output, indent=2), encoding="utf-8")
    print(json.dumps(output, indent=2))
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
