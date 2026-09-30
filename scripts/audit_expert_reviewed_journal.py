from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader


def metrics(tp, tn, fp, fn, expected):
    n = tp + tn + fp + fn
    p = 100 * tp / (tp + fp) if tp + fp else 0
    r = 100 * tp / (tp + fn) if tp + fn else 0
    values = [100 * (tp + tn) / n, p, r, 2 * p * r / (p + r) if p + r else 0]
    return all(math.isclose(a, b, abs_tol=0.011) for a, b in zip(values, expected))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx", type=Path)
    ap.add_argument("pdf", type=Path)
    ap.add_argument("change_report", type=Path)
    ap.add_argument("output", type=Path)
    a = ap.parse_args()
    doc = Document(a.docx)
    text = "\n".join(p.text for p in doc.paragraphs)
    tables = "\n".join(cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    all_text = (text + "\n" + tables).lower()
    report_in = json.loads(a.change_report.read_text(encoding="utf-8"))
    refs = [p.text for p in doc.paragraphs if re.match(r"^\[\d+\]", p.text.strip())]
    nums = [int(re.match(r"^\[(\d+)\]", r.strip()).group(1)) for r in refs]
    banned = ["grade c", "graded c", "single-reviewer", "single-rater", "one administrator", "n=30", "n = 30", "30 cases", "30/model", "administrator-verified"]
    pdf = PdfReader(str(a.pdf))
    blank = [i+1 for i,p in enumerate(pdf.pages) if not (p.extract_text() or "").strip()]
    with ZipFile(a.docx) as z:
        crc = z.testzip()
        media_count = len([n for n in z.namelist() if n.startswith("word/media/")])
    metric_ok = all([
        metrics(1,19,0,10,[66.67,100,9.09,16.67]),
        metrics(12,16,2,0,[93.33,85.71,100,92.31]),
        metrics(27,0,3,0,[90,90,100,94.74]),
        metrics(206,1262,63,29,[94.10,76.58,87.66,81.75]),
        metrics(150,1296,30,84,[92.69,83.33,64.10,72.46]),
        metrics(177,1261,65,57,[92.18,73.14,75.64,74.37]),
    ])
    out = {
        "sha256": hashlib.sha256(a.docx.read_bytes()).hexdigest(),
        "zipCrcFailure": crc,
        "paragraphs": len(doc.paragraphs), "tables": len(doc.tables), "inlineShapes": len(doc.inline_shapes), "mediaParts": media_count,
        "referenceCount": len(refs), "referencesSequential": nums == list(range(1,115)),
        "referencesUnchanged": report_in["referencesUnchanged"], "citationSequencePreserved": report_in["citationSequencePreserved"],
        "bannedWordingHits": {x: all_text.count(x) for x in banned},
        "panelReviewPresent": "a panel of domain experts" in all_text,
        "agreementPresent": "agreed with the final judgments" in all_text,
        "reliabilityNotInvented": "no inter-rater reliability statistic is claimed" in all_text,
        "allMetricsRecomputed": metric_ok,
        "pdfPages": len(pdf.pages), "blankPdfPages": blank,
    }
    out["passed"] = all([crc is None, len(doc.paragraphs)==299, len(doc.tables)==24, len(doc.inline_shapes)==15, media_count==15, len(refs)==114, out["referencesSequential"], out["referencesUnchanged"], out["citationSequencePreserved"], not any(out["bannedWordingHits"].values()), out["panelReviewPresent"], out["agreementPresent"], out["reliabilityNotInvented"], metric_ok, len(pdf.pages)==20, not blank])
    a.output.write_text(json.dumps(out, indent=2), encoding="utf-8")
    print(json.dumps(out, indent=2))
    if not out["passed"]: raise SystemExit(1)


if __name__ == "__main__": main()
