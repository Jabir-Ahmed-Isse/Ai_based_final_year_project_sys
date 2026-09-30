from __future__ import annotations

import json
import re
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from lxml import etree

from build_research_report import INK, add_paragraph, set_font


ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "outputs" / "journal-package-20260724"
SOURCE = PACKAGE / "final-journal-manuscript-complete-visuals.docx"
OUTPUT = PACKAGE / "final-journal-manuscript-symbolic-mathematics.docx"
REPORT = PACKAGE / "symbolic-mathematics-audit.json"
MML2OMML = Path(r"C:\Program Files\Microsoft Office\Office16\MML2OMML.XSL")


MATH_NS = "http://www.w3.org/1998/Math/MathML"
OMML_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def math(body: str) -> str:
    return f'<math xmlns="{MATH_NS}"><mrow>{body}</mrow></math>'


def fn(name: str, argument: str) -> str:
    return (
        f'<mi mathvariant="normal">{name}</mi><mo>(</mo>{argument}<mo>)</mo>'
    )


def sub(base: str, index: str) -> str:
    return f"<msub><mrow>{base}</mrow><mrow>{index}</mrow></msub>"


def sup(base: str, exponent: str) -> str:
    return f"<msup><mrow>{base}</mrow><mrow>{exponent}</mrow></msup>"


def frac(numerator: str, denominator: str) -> str:
    return f"<mfrac><mrow>{numerator}</mrow><mrow>{denominator}</mrow></mfrac>"


def summation(index: str, upper: str, expression: str) -> str:
    return (
        "<munderover><mo>∑</mo>"
        f"<mrow>{index}</mrow><mrow>{upper}</mrow></munderover>"
        f"<mrow>{expression}</mrow>"
    )


def norm(expression: str) -> str:
    return sub(
        f"<mrow><mo>∥</mo>{expression}<mo>∥</mo></mrow>",
        "<mn>2</mn>",
    )


def metric(name: str) -> str:
    return f'<mi mathvariant="normal">{name}</mi>'


def move_before(paragraph, target) -> None:
    target._p.addprevious(paragraph._p)


def insert_heading(document: Document, target, text: str, level: int = 3) -> None:
    paragraph = document.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    move_before(paragraph, target)


def insert_text(document: Document, target, text: str) -> None:
    paragraph = add_paragraph(document, text)
    move_before(paragraph, target)


def insert_definition(document: Document, target, label: str, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(18)
    paragraph.paragraph_format.space_after = Pt(4)
    label_run = paragraph.add_run(label)
    set_font(label_run, bold=True, color=INK)
    body_run = paragraph.add_run(text)
    set_font(body_run, color=INK)
    move_before(paragraph, target)


def build_transform() -> etree.XSLT:
    return etree.XSLT(etree.parse(str(MML2OMML)))


def insert_equation(
    document: Document,
    target,
    transformer: etree.XSLT,
    mathml: str,
) -> None:
    source = etree.fromstring(mathml.encode("utf-8"))
    transformed = transformer(source)
    omml_root = deepcopy(transformed.getroot())

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(7)

    math_para = etree.Element(f"{{{OMML_NS}}}oMathPara", nsmap={"m": OMML_NS})
    math_para.append(omml_root)
    paragraph._p.append(math_para)
    move_before(paragraph, target)


def replace_plain_formula_text(document: Document) -> int:
    replacements = {
        "C(78,2)=3,003 unique pairs per model. The supervisor design yields 78x20=1,560 candidates per model.": (
            "3,003 unique project pairs per model and 1,560 project-supervisor candidates "
            "per model. Both totals are derived symbolically in Section 3.11."
        ),
        "The number of unordered pairs was independently recomputed as n(n−1)/2, yielding 3,003 for n=78.": (
            "The number of unordered pairs was independently recomputed using the "
            "combinatorial expression in Section 3.11, yielding 3,003 for 78 projects."
        ),
        "rho=0.954": "ρₛ = 0.954",
        "rho=0.9537": "ρₛ = 0.9537",
    }
    changed = 0
    for paragraph in document.paragraphs:
        original = paragraph.text
        revised = original
        for old, new in replacements.items():
            revised = revised.replace(old, new)
        if revised != original:
            paragraph.text = revised
            changed += 1
    return changed


def equations() -> list[tuple[str, list[tuple[str, str]], list[tuple[str, str]]]]:
    tf = fn("tf", "<mi>t</mi><mo>,</mo><mi>d</mi>")
    idf = fn("idf", "<mi>t</mi>")
    df = fn("df", "<mi>t</mi>")
    f_td = fn("f", "<mi>t</mi><mo>,</mo><mi>d</mi>")

    return [
        (
            "3.11.1 Corpus size and experimental cardinality",
            [
                (
                    "corpus",
                    math(
                        "<mi>D</mi><mo>=</mo><mo>{</mo>"
                        f"{sub('<mi>d</mi>', '<mn>1</mn>')}<mo>,</mo>"
                        f"{sub('<mi>d</mi>', '<mn>2</mn>')}<mo>,</mo><mo>…</mo><mo>,</mo>"
                        f"{sub('<mi>d</mi>', '<mi>N</mi>')}<mo>}}</mo>"
                        "<mo>,</mo><mspace width=\"0.5em\"/>"
                        "<mo>|</mo><mi>D</mi><mo>|</mo><mo>=</mo><mi>N</mi>"
                    ),
                ),
                (
                    "pairs",
                    math(
                        f"{sub('<mi>N</mi>', '<mtext>pair</mtext>')}<mo>=</mo>"
                        f"{frac('<mi>n</mi><mo>(</mo><mi>n</mi><mo>−</mo><mn>1</mn><mo>)</mo>', '<mn>2</mn>')}"
                        "<mo>=</mo>"
                        f"{frac('<mn>78</mn><mo>×</mo><mn>77</mn>', '<mn>2</mn>')}"
                        "<mo>=</mo><mn>3003</mn>"
                    ),
                ),
                (
                    "candidates",
                    math(
                        f"{sub('<mi>N</mi>', '<mtext>cand</mtext>')}<mo>=</mo>"
                        f"{sub('<mi>n</mi>', '<mi>p</mi>')}<mo>×</mo>"
                        f"{sub('<mi>n</mi>', '<mi>s</mi>')}<mo>=</mo>"
                        "<mn>78</mn><mo>×</mo><mn>20</mn><mo>=</mo><mn>1560</mn>"
                    ),
                ),
            ],
            [
                ("D, N. ", "Project corpus and its cardinality."),
                ("nₚ, nₛ. ", "Numbers of projects and supervisor profiles."),
            ],
        ),
        (
            "3.11.2 TF-IDF and vector similarity",
            [
                (
                    "tf",
                    math(
                        f"{tf}<mo>=</mo>"
                        f"{frac(f_td, summation('<mi>u</mi><mo>∈</mo><mi>d</mi>', '', fn('f', '<mi>u</mi><mo>,</mo><mi>d</mi>')))}"
                    ),
                ),
                (
                    "idf",
                    math(
                        f"{idf}<mo>=</mo><mi mathvariant=\"normal\">log</mi>"
                        "<mo>(</mo>"
                        f"{frac('<mi>N</mi><mo>+</mo><mn>1</mn>', f'{df}<mo>+</mo><mn>1</mn>')}"
                        "<mo>)</mo><mo>+</mo><mn>1</mn>"
                    ),
                ),
                (
                    "weight",
                    math(
                        f"{fn('w', '<mi>t</mi><mo>,</mo><mi>d</mi>')}<mo>=</mo>"
                        f"{tf}<mo>×</mo>{idf}"
                    ),
                ),
                (
                    "cosine",
                    math(
                        f"{fn('sim', '<mi>x</mi><mo>,</mo><mi>y</mi>')}<mo>=</mo>"
                        f"{frac(f'{sup('<mi>x</mi>', '<mi>T</mi>')}<mi>y</mi>', norm('<mi>x</mi>') + norm('<mi>y</mi>'))}"
                    ),
                ),
            ],
            [
                ("f(t,d). ", "Frequency of term t in document d."),
                ("df(t). ", "Number of documents containing t."),
                ("x, y. ", "Sparse or dense vector representations compared by cosine similarity."),
            ],
        ),
        (
            "3.11.3 Transformer representation and multi-field aggregation",
            [
                (
                    "transformer",
                    math(
                        "<mi>H</mi><mo>=</mo>"
                        f"{fn('Transformer', f'{sub('<mi>x</mi>', '<mn>1</mn>')}<mo>,</mo><mo>…</mo><mo>,</mo>{sub('<mi>x</mi>', '<mi>n</mi>')}')}"
                    ),
                ),
                (
                    "embedding",
                    math(
                        f"{fn('e', '<mi>d</mi>')}<mo>=</mo>"
                        f"{frac(fn('Pool', '<mi>H</mi>'), norm(fn('Pool', '<mi>H</mi>')))}"
                    ),
                ),
                (
                    "multifield",
                    math(
                        f"{fn('S', f'{sub('<mi>p</mi>', '<mi>i</mi>')}<mo>,</mo>{sub('<mi>p</mi>', '<mi>j</mi>')}')}<mo>=</mo>"
                        f"{frac(summation('<mi>f</mi><mo>∈</mo><mi>F</mi>', '', f'{sub('<mi>w</mi>', '<mi>f</mi>')}{sub('<mi>s</mi>', '<mi>f</mi>')}<mo>(</mo>{sub('<mi>p</mi>', '<mi>i</mi>')}<mo>,</mo>{sub('<mi>p</mi>', '<mi>j</mi>')}<mo>)</mo>{sub('<mi>I</mi>', '<mi>f</mi>')}'), summation('<mi>f</mi><mo>∈</mo><mi>F</mi>', '', f'{sub('<mi>w</mi>', '<mi>f</mi>')}{sub('<mi>I</mi>', '<mi>f</mi>')}'))}"
                    ),
                ),
            ],
            [
                ("H. ", "Contextual token-state matrix."),
                ("e(d). ", "L2-normalized sentence or document embedding."),
                ("F, w_f, s_f. ", "Available fields, fixed field weights and field-level similarity."),
                ("I_f. ", "Availability indicator: 1 when field f exists for both projects, otherwise 0."),
            ],
        ),
        (
            "3.11.4 Risk bands and workload-aware supervisor ranking",
            [
                (
                    "risk",
                    math(
                        f"{fn('Risk', '<mi>S</mi>')}<mo>=</mo><mo>{{</mo>"
                        "<mtable columnalign=\"left left\" rowspacing=\"0.2em\">"
                        "<mtr><mtd><mtext>Low</mtext></mtd><mtd><mi>S</mi><mo>&lt;</mo><mn>40</mn></mtd></mtr>"
                        "<mtr><mtd><mtext>Medium</mtext></mtd><mtd><mn>40</mn><mo>≤</mo><mi>S</mi><mo>&lt;</mo><mn>70</mn></mtd></mtr>"
                        "<mtr><mtd><mtext>High</mtext></mtd><mtd><mi>S</mi><mo>≥</mo><mn>70</mn></mtd></mtr>"
                        "</mtable>"
                    ),
                ),
                (
                    "semantic",
                    math(
                        f"{fn('Q', '<mi>p</mi><mo>,</mo><mi>s</mi>')}<mo>=</mo>"
                        f"{summation('<mi>k</mi><mo>∈</mo><mi>K</mi>', '', f'{sub('<mi>α</mi>', '<mi>k</mi>')}{sub('<mi>q</mi>', '<mi>k</mi>')}<mo>(</mo><mi>p</mi><mo>,</mo><mi>s</mi><mo>)</mo>')}"
                        "<mo>,</mo><mspace width=\"0.5em\"/>"
                        f"{summation('<mi>k</mi><mo>∈</mo><mi>K</mi>', '', sub('<mi>α</mi>', '<mi>k</mi>'))}<mo>=</mo><mn>1</mn>"
                    ),
                ),
                (
                    "utilization",
                    math(
                        f"{fn('u', '<mi>s</mi>')}<mo>=</mo>"
                        f"{frac(sub('<mi>c</mi>', '<mi>s</mi>'), sub('<mi>C</mi>', '<mi>s</mi>'))}"
                        "<mo>,</mo><mspace width=\"0.5em\"/><mn>0</mn><mo>≤</mo>"
                        f"{fn('u', '<mi>s</mi>')}<mo>≤</mo><mn>1</mn>"
                    ),
                ),
                (
                    "penalty",
                    math(
                        f"{fn('P', '<mi>s</mi>')}<mo>=</mo>"
                        f"{fn('min', '<mn>15</mn><mo>,</mo><mi>λ</mi>' + fn('u', '<mi>s</mi>'))}"
                    ),
                ),
                (
                    "adjusted",
                    math(
                        f"{fn('R', '<mi>p</mi><mo>,</mo><mi>s</mi>')}<mo>=</mo>"
                        f"{fn('Q', '<mi>p</mi><mo>,</mo><mi>s</mi>')}<mo>−</mo>"
                        f"{fn('P', '<mi>s</mi>')}"
                    ),
                ),
                (
                    "assignment",
                    math(
                        f"{sup('<mi>s</mi>', '<mo>*</mo>')}<mo>(</mo><mi>p</mi><mo>)</mo><mo>=</mo>"
                        "<munder><mi mathvariant=\"normal\">arg max</mi>"
                        "<mrow><mi>s</mi><mo>∈</mo><mi>E</mi><mo>(</mo><mi>p</mi><mo>)</mo>"
                        "<mo>,</mo><msub><mi>c</mi><mi>s</mi></msub><mo>&lt;</mo>"
                        "<msub><mi>C</mi><mi>s</mi></msub></mrow></munder>"
                        f"{fn('R', '<mi>p</mi><mo>,</mo><mi>s</mi>')}"
                    ),
                ),
            ],
            [
                ("Q(p,s). ", "Pure semantic supervisor score."),
                ("c_s, C_s. ", "Current load and maximum capacity of supervisor s."),
                ("E(p). ", "Eligible supervisors for project p."),
                ("λ. ", "Utilization-penalty scale; the persisted implementation caps the penalty at 15 points."),
            ],
        ),
        (
            "3.11.5 Descriptive and agreement statistics",
            [
                (
                    "mean",
                    math(
                        f"{sub('<mover><mi>x</mi><mo>¯</mo></mover>', '<mi>m</mi>')}<mo>=</mo>"
                        f"{frac('<mn>1</mn>', '<mi>n</mi>')}"
                        f"{summation('<mi>i</mi><mo>=</mo><mn>1</mn>', '<mi>n</mi>', sub('<mi>x</mi>', '<mi>i</mi><mi>m</mi>'))}"
                    ),
                ),
                (
                    "sd",
                    math(
                        f"{sub('<mi>s</mi>', '<mi>m</mi>')}<mo>=</mo><msqrt>"
                        f"{frac(summation('<mi>i</mi><mo>=</mo><mn>1</mn>', '<mi>n</mi>', f'<msup><mrow>{sub('<mi>x</mi>', '<mi>i</mi><mi>m</mi>')}<mo>−</mo>{sub('<mover><mi>x</mi><mo>¯</mo></mover>', '<mi>m</mi>')}</mrow><mn>2</mn></msup>'), '<mi>n</mi><mo>−</mo><mn>1</mn>')}"
                        "</msqrt>"
                    ),
                ),
                (
                    "ci",
                    math(
                        f"{sub('<mover><mi>x</mi><mo>¯</mo></mover>', '<mi>m</mi>')}"
                        "<mo>±</mo><mn>1.96</mn>"
                        f"{frac(sub('<mi>s</mi>', '<mi>m</mi>'), '<msqrt><mi>n</mi></msqrt>')}"
                    ),
                ),
                (
                    "pearson",
                    math(
                        f"{sub('<mi>r</mi>', '<mi>x</mi><mi>y</mi>')}<mo>=</mo>"
                        f"{frac(summation('<mi>i</mi><mo>=</mo><mn>1</mn>', '<mi>n</mi>', f'<mrow><mo>(</mo>{sub('<mi>x</mi>', '<mi>i</mi>')}<mo>−</mo><mover><mi>x</mi><mo>¯</mo></mover><mo>)</mo><mo>(</mo>{sub('<mi>y</mi>', '<mi>i</mi>')}<mo>−</mo><mover><mi>y</mi><mo>¯</mo></mover><mo>)</mo></mrow>'), f'<msqrt>{summation('<mi>i</mi><mo>=</mo><mn>1</mn>', '<mi>n</mi>', f'<msup><mrow><mo>(</mo>{sub('<mi>x</mi>', '<mi>i</mi>')}<mo>−</mo><mover><mi>x</mi><mo>¯</mo></mover><mo>)</mo></mrow><mn>2</mn></msup>')}{summation('<mi>i</mi><mo>=</mo><mn>1</mn>', '<mi>n</mi>', f'<msup><mrow><mo>(</mo>{sub('<mi>y</mi>', '<mi>i</mi>')}<mo>−</mo><mover><mi>y</mi><mo>¯</mo></mover><mo>)</mo></mrow><mn>2</mn></msup>')}</msqrt>')}"
                    ),
                ),
                (
                    "spearman",
                    math(
                        f"{sub('<mi>ρ</mi>', '<mi>s</mi>')}<mo>=</mo>"
                        f"{fn('r', f'{fn('rank', '<mi>x</mi>')}<mo>,</mo>{fn('rank', '<mi>y</mi>')}')}"
                    ),
                ),
                (
                    "top_agreement",
                    math(
                        f"{sub('<mi>A</mi>', '<mtext>top</mtext>')}<mo>=</mo>"
                        f"{frac('<mn>1</mn>', '<mi>n</mi>')}"
                        f"{summation('<mi>i</mi><mo>=</mo><mn>1</mn>', '<mi>n</mi>', '<mi mathvariant=\"double-struck\">I</mi><mo>[</mo>' + sup(sub('<mi>s</mi>', '<mi>i</mi>'), '<mo>(</mo><mi>a</mi><mo>)</mo>') + '<mo>=</mo>' + sup(sub('<mi>s</mi>', '<mi>i</mi>'), '<mo>(</mo><mi>b</mi><mo>)</mo>') + '<mo>]</mo>')}"
                    ),
                ),
            ],
            [
                ("x_im. ", "Score for item i under model m."),
                ("r_xy, ρ_s. ", "Pearson and Spearman association coefficients."),
                ("A_top. ", "Proportion of projects for which two models select the same top supervisor."),
            ],
        ),
        (
            "3.11.6 Human-label-dependent evaluation metrics",
            [
                (
                    "accuracy",
                    math(
                        f"{metric('Accuracy')}<mo>=</mo>"
                        f"{frac('<mi>TP</mi><mo>+</mo><mi>TN</mi>', '<mi>TP</mi><mo>+</mo><mi>TN</mi><mo>+</mo><mi>FP</mi><mo>+</mo><mi>FN</mi>')}"
                    ),
                ),
                (
                    "precision_recall",
                    math(
                        f"{metric('Precision')}<mo>=</mo>"
                        f"{frac('<mi>TP</mi>', '<mi>TP</mi><mo>+</mo><mi>FP</mi>')}"
                        "<mo>,</mo><mspace width=\"1em\"/>"
                        f"{metric('Recall')}<mo>=</mo>"
                        f"{frac('<mi>TP</mi>', '<mi>TP</mi><mo>+</mo><mi>FN</mi>')}"
                    ),
                ),
                (
                    "f1",
                    math(
                        f"{sub('<mi>F</mi>', '<mn>1</mn>')}<mo>=</mo>"
                        f"{frac('<mn>2</mn><mo>×</mo>' + metric('Precision') + '<mo>×</mo>' + metric('Recall'), metric('Precision') + '<mo>+</mo>' + metric('Recall'))}"
                    ),
                ),
                (
                    "mcc",
                    math(
                        "<mi>MCC</mi><mo>=</mo>"
                        f"{frac('<mi>TP</mi><mo>×</mo><mi>TN</mi><mo>−</mo><mi>FP</mi><mo>×</mo><mi>FN</mi>', '<msqrt><mrow><mo>(</mo><mi>TP</mi><mo>+</mo><mi>FP</mi><mo>)</mo><mo>(</mo><mi>TP</mi><mo>+</mo><mi>FN</mi><mo>)</mo><mo>(</mo><mi>TN</mi><mo>+</mo><mi>FP</mi><mo>)</mo><mo>(</mo><mi>TN</mi><mo>+</mo><mi>FN</mi><mo>)</mo></mrow></msqrt>')}"
                    ),
                ),
                (
                    "mae_rmse",
                    math(
                        "<mi>MAE</mi><mo>=</mo>"
                        f"{frac('<mn>1</mn>', '<mi>n</mi>')}"
                        f"{summation('<mi>i</mi><mo>=</mo><mn>1</mn>', '<mi>n</mi>', '<mo>|</mo>' + sub('<mi>y</mi>', '<mi>i</mi>') + '<mo>−</mo>' + sub('<mover><mi>y</mi><mo>^</mo></mover>', '<mi>i</mi>') + '<mo>|</mo>')}"
                        "<mo>,</mo><mspace width=\"1em\"/><mi>RMSE</mi><mo>=</mo><msqrt>"
                        f"{frac(summation('<mi>i</mi><mo>=</mo><mn>1</mn>', '<mi>n</mi>', '<msup><mrow>' + sub('<mi>y</mi>', '<mi>i</mi>') + '<mo>−</mo>' + sub('<mover><mi>y</mi><mo>^</mo></mover>', '<mi>i</mi>') + '</mrow><mn>2</mn></msup>'), '<mi>n</mi>')}"
                        "</msqrt>"
                    ),
                ),
                (
                    "mrr",
                    math(
                        "<mi>MRR</mi><mo>=</mo>"
                        f"{frac('<mn>1</mn>', '<mo>|</mo><mi>Q</mi><mo>|</mo>')}"
                        f"{summation('<mi>q</mi><mo>∈</mo><mi>Q</mi>', '', frac('<mn>1</mn>', sub('<mi>rank</mi>', '<mi>q</mi>')))}"
                    ),
                ),
                (
                    "map",
                    math(
                        "<mi>MAP</mi><mo>=</mo>"
                        f"{frac('<mn>1</mn>', '<mo>|</mo><mi>Q</mi><mo>|</mo>')}"
                        f"{summation('<mi>q</mi><mo>∈</mo><mi>Q</mi>', '', sub('<mi>AP</mi>', '<mi>q</mi>'))}"
                    ),
                ),
                (
                    "ndcg",
                    math(
                        f"{sub('<mi>NDCG</mi>', '<mi>K</mi>')}<mo>=</mo>"
                        f"{frac(sub('<mi>DCG</mi>', '<mi>K</mi>'), sub('<mi>IDCG</mi>', '<mi>K</mi>'))}"
                        "<mo>,</mo><mspace width=\"1em\"/>"
                        f"{sub('<mi>DCG</mi>', '<mi>K</mi>')}<mo>=</mo>"
                        f"{summation('<mi>k</mi><mo>=</mo><mn>1</mn>', '<mi>K</mi>', frac('<msup><mn>2</mn>' + sub('<mi>rel</mi>', '<mi>k</mi>') + '</msup><mo>−</mo><mn>1</mn>', '<msub><mi mathvariant=\"normal\">log</mi><mn>2</mn></msub><mo>(</mo><mi>k</mi><mo>+</mo><mn>1</mn><mo>)</mo>'))}"
                    ),
                ),
            ],
            [
                ("TP, TN, FP, FN. ", "Confusion-matrix counts at a frozen decision threshold."),
                ("Q. ", "Set of evaluation queries or projects with expert relevance judgements."),
                ("Important validity boundary. ", "These equations define the prespecified evaluation only. With zero expert-consensus labels, the journal does not report numerical accuracy, F1, MCC, MAE, RMSE, MRR, MAP or NDCG results."),
            ],
        ),
    ]


def main() -> None:
    if not MML2OMML.exists():
        raise FileNotFoundError(MML2OMML)

    document = Document(SOURCE)
    plain_text_changes = replace_plain_formula_text(document)

    start_matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip() == "3.11 Mathematical formulation"
    ]
    end_matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip() == "3.12 Implementation architecture"
    ]
    if len(start_matches) != 1 or len(end_matches) != 1:
        raise RuntimeError("Could not uniquely locate Sections 3.11 and 3.12.")
    start = start_matches[0]
    target = end_matches[0]

    paragraphs = document.paragraphs
    start_index = next(index for index, paragraph in enumerate(paragraphs) if paragraph._p is start._p)
    end_index = next(index for index, paragraph in enumerate(paragraphs) if paragraph._p is target._p)
    removed = 0
    for paragraph in paragraphs[start_index + 1 : end_index]:
        parent = paragraph._element.getparent()
        parent.remove(paragraph._element)
        removed += 1

    transformer = build_transform()
    insert_text(
        document,
        target,
        "The implemented calculations are stated below using the same symbolic convention "
        "as the thesis. Equations define the operational pipeline and the prespecified "
        "evaluation; they do not convert unavailable human-label metrics into reported results.",
    )

    equation_count = 0
    for heading, equation_items, definitions in equations():
        insert_heading(document, target, heading, level=3)
        for _, equation_mathml in equation_items:
            insert_equation(document, target, transformer, equation_mathml)
            equation_count += 1
        for label, text in definitions:
            insert_definition(document, target, label, text)

    document.save(OUTPUT)

    reopened = Document(OUTPUT)
    paragraph_text = [paragraph.text.strip() for paragraph in reopened.paragraphs]
    with zipfile.ZipFile(OUTPUT) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        equation_objects = len(re.findall(r"<m:oMath\b", xml))
        equation_paragraphs = len(re.findall(r"<m:oMathPara\b", xml))
        fractions = len(re.findall(r"<m:f\b", xml))
        radicals = len(re.findall(r"<m:rad\b", xml))
        subscripts = len(re.findall(r"<m:sSub\b", xml))
        superscripts = len(re.findall(r"<m:sSup\b", xml))
        matrices = len(re.findall(r"<m:m\b", xml))

    report = {
        "source": str(SOURCE.resolve()),
        "output": str(OUTPUT.resolve()),
        "removed_plain_formula_list_paragraphs": removed,
        "plain_text_symbol_replacements": plain_text_changes,
        "inserted_display_equations": equation_count,
        "omml_equation_objects": equation_objects,
        "omml_equation_paragraphs": equation_paragraphs,
        "omml_fractions": fractions,
        "omml_radicals": radicals,
        "omml_subscripts": subscripts,
        "omml_superscripts": superscripts,
        "omml_matrices": matrices,
        "embedded_media_preserved": len(media),
        "visualization_results_section_preserved": (
            "4.8 Live visualization-page analysis" in paragraph_text
        ),
        "comparative_evaluation_preserved": (
            "4.4 Comparative evaluation and evidence-bounded rankings" in paragraph_text
        ),
        "reference_section_preserved": "References" in paragraph_text,
    }
    REPORT.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
