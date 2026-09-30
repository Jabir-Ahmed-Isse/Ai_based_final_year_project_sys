from __future__ import annotations

import csv
import hashlib
import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader


DOWNLOADS = Path(r"C:\Users\YoGa\Downloads")
THESIS = Path(r"M:\THESIS_BOOK (Autosaved).docx")
OUT = Path(r"M:\hormuud-academic-project-system\outputs\journal-package-20260724\corpus")
OUT.mkdir(parents=True, exist_ok=True)

DATE_PREFIX = "2026-07-24"
doi_re = re.compile(r"\b10\.\d{4,9}/[-._;()/:A-Z0-9]+\b", re.I)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


records: list[dict[str, object]] = []
full_text: dict[str, str] = {}

for path in sorted(DOWNLOADS.iterdir(), key=lambda p: p.name.casefold()):
    if not path.is_file() or path.suffix.lower() not in {".pdf", ".doc", ".zip"}:
        continue
    if path.stat().st_mtime < 1784840000:  # focus on the supplied 24 July corpus
        continue
    record: dict[str, object] = {
        "file": path.name,
        "path": str(path),
        "extension": path.suffix.lower(),
        "bytes": path.stat().st_size,
        "sha256": sha256(path),
        "pages": "",
        "document_title": "",
        "document_author": "",
        "subject": "",
        "doi_candidates": "",
        "extraction_status": "",
        "first_pages_excerpt": "",
    }
    try:
        if path.suffix.lower() == ".pdf":
            reader = PdfReader(str(path))
            record["pages"] = len(reader.pages)
            meta = reader.metadata or {}
            record["document_title"] = clean(meta.get("/Title"))
            record["document_author"] = clean(meta.get("/Author"))
            record["subject"] = clean(meta.get("/Subject"))
            text_parts = []
            for page in reader.pages[:3]:
                text_parts.append(page.extract_text() or "")
            text = clean("\n".join(text_parts))
            record["doi_candidates"] = "; ".join(sorted(set(m.rstrip(".,;)") for m in doi_re.findall(text))))
            record["first_pages_excerpt"] = text[:5000]
            record["extraction_status"] = "ok" if text else "no extractable text"
            full_text[path.name] = text
        elif path.suffix.lower() == ".zip":
            with zipfile.ZipFile(path) as archive:
                names = archive.namelist()
            record["document_title"] = "Archive contents"
            record["first_pages_excerpt"] = " | ".join(names[:200])
            record["extraction_status"] = f"archive listed ({len(names)} entries)"
        else:
            record["extraction_status"] = "legacy .doc; metadata extraction pending conversion"
    except Exception as exc:
        record["extraction_status"] = f"error: {type(exc).__name__}: {exc}"
    records.append(record)

# Thesis extraction, including tables because requirements/results often live there.
thesis_doc = Document(str(THESIS))
thesis_lines: list[str] = []
for para in thesis_doc.paragraphs:
    text = clean(para.text)
    if text:
        thesis_lines.append(text)
for table_index, table in enumerate(thesis_doc.tables, 1):
    thesis_lines.append(f"[TABLE {table_index}]")
    for row in table.rows:
        thesis_lines.append(" | ".join(clean(cell.text) for cell in row.cells))
thesis_text = "\n".join(thesis_lines)
(OUT / "thesis-extracted-text.txt").write_text(thesis_text, encoding="utf-8")

with (OUT / "paper-corpus-inventory.csv").open("w", newline="", encoding="utf-8-sig") as stream:
    writer = csv.DictWriter(stream, fieldnames=list(records[0].keys()))
    writer.writeheader()
    writer.writerows(records)

(OUT / "paper-first-pages.json").write_text(
    json.dumps(full_text, ensure_ascii=False, indent=2), encoding="utf-8"
)
(OUT / "corpus-summary.json").write_text(
    json.dumps(
        {
            "paper_files": len(records),
            "unique_hashes": len({r["sha256"] for r in records}),
            "duplicate_files": len(records) - len({r["sha256"] for r in records}),
            "pdf_files": sum(r["extension"] == ".pdf" for r in records),
            "pdfs_with_text": sum(r["extension"] == ".pdf" and r["extraction_status"] == "ok" for r in records),
            "doi_candidates_found": sum(bool(r["doi_candidates"]) for r in records),
            "thesis_paragraphs": len(thesis_doc.paragraphs),
            "thesis_tables": len(thesis_doc.tables),
            "thesis_characters": len(thesis_text),
        },
        indent=2,
    ),
    encoding="utf-8",
)
print((OUT / "corpus-summary.json").read_text(encoding="utf-8"))
