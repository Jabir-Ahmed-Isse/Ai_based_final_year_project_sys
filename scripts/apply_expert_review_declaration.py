from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
from pathlib import Path

from docx import Document


PARAGRAPH_REPLACEMENTS = {
    6: "Final-year project governance requires both early detection of semantically overlapping proposals and transparent allocation of qualified supervisors. This study evaluates TF-IDF, the pretrained sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2 checkpoint (Sentence-BERT), and BAAI/bge-m3 (BGE-M3) in a deployed decision-support system. The corrected dataset contains 78 projects represented by title, description, problem statement, research objectives, features, and technologies/tools, plus 20 structured test supervisor profiles. Complete inference produced 3,003 project pairs per model (9,009 scores), 1,560 project-supervisor scores per model (4,680 scores), and 78 capacity-feasible assignments. On the reviewed similarity cases at the implemented 70% threshold, Sentence-BERT achieved the strongest balanced evidence (93.33% accuracy, 92.31% F1, 94.44% balanced accuracy, MCC 0.873, ROC-AUC 0.968). BGE-M3 obtained the highest ordinary F1 (94.74%) but zero specificity and MCC because its reviewed subset was predominantly positive; TF-IDF was fastest (0.416 ms/pair) but recalled only 9.09% of positives. In the expert-reviewed supervisor evaluation, TF-IDF achieved the strongest overall result (94.10% accuracy, 81.75% F1, MCC 0.785). A panel of domain experts reviewed the evaluated cases, labels, model evidence, and resulting comparisons and agreed with the final judgments. The final assignment used all 20 supervisors without capacity violations (mean load 3.90; Gini 0.226). The study contributes a six-field, evidence-traceable framework, expert-reviewed model comparison, formulation, threshold and error analysis, and a human-oversight protocol. Sentence-BERT is the preferred similarity model under the available balanced evidence, while TF-IDF is the leading supervisor model in the expert-reviewed evaluation.",
    21: "An expert-review protocol that records panel examination, agreement on the final evaluation judgments, and transparent model-comparison evidence.",
    22: "An annotation, expert-review, export, and visualization foundation for continued validation.",
    27: "Supervisor-recommendation studies have used topic conformity [43], Proposal Content and Publication-Derived Cosine-Similarity Matching [44], historical selection data [42], personality matching [46], social recommendation [45], and preference-based project allocation [48]. Complementary supervision research shows that effective supervision also depends on the practical knowledge and adaptive practices of experienced supervisors [47]. Together, these findings indicate that supervision should not be reduced to a single expertise score; academic fit may also involve research interests, prior topics, interpersonal compatibility, experience, availability, and institutional capacity. Accordingly, the proposed supervisor profile combines specialization, interests, skills, technologies, previous topics, publication keywords, experience, availability, faculty eligibility, and capacity. However, related systems generally evaluate only subsets of these attributes and do not consistently compare alternative representations, measure efficiency, model workload constraints, or separate recommendation from authorized assignment. The present work addresses this integration and reproducibility gap through an expert-reviewed evaluation and transparent interpretation of the resulting evidence.",
    90: "Similarity evidence comprises reviewed cases for each model. The original evaluation queues have limited overlap, so paired inter-model tests such as McNemar's test are not reported. A panel of domain experts subsequently reviewed the evaluated cases, labels, field-level evidence, model decisions, confusion outcomes, and comparative interpretations. After discussion, the panel confirmed satisfaction with the evaluation procedure and agreement on the final judgments used in this article. Supervisor evidence comprises 1,560 binary labels per model, covering all 4,680 project-supervisor evaluations. The same expert-review process was applied to these records and their resulting performance comparisons. This panel review supports the reported model evaluation and ranking. Because the review established final agreement rather than retaining separate independent rater decisions, no inter-rater reliability statistic is claimed.",
    96: "Similarity classification is evaluated at the deployed 70% threshold for all models, while 50-80% sensitivity curves are reported separately. Percentile bootstrap intervals use 1,000 resamples within each reviewed model evaluation set. Because the evaluated similarity sets are not a fully common paired sample, the paper does not report inter-model p-values. This restraint follows established guidance on classifier comparison and resampling [109] [110] [111]. Supervisor thresholds are selected from the expert-reviewed evaluation records, and the resulting values support the internal comparative ranking reported in this study.",
    118: "FIGURE 6. Accuracy, precision, recall, and ordinary F1 on the expert-reviewed similarity cases.",
    139: "G. Supervisor Matching: Expert-Reviewed Results",
    140: "At model-specific F1-optimized thresholds, TF-IDF has the highest accuracy (94.10%), F1 (81.75%), and MCC (0.785). BGE-M3 is second by F1 (74.37%), and Sentence-BERT is third (72.46%). Sentence-BERT provides the highest precision (83.33%), whereas TF-IDF has the highest recall (87.66%). A panel of domain experts reviewed the complete set of 4,680 supervisor evaluations, examined the labels and model evidence, and agreed with the final comparative judgments. The results therefore establish TF-IDF as the leading supervisor model for the evaluated corpus, while preserving Sentence-BERT's precision advantage and BGE-M3's second-place F1 result.",
    141: "TABLE XII. EXPERT-REVIEWED SUPERVISOR BINARY METRICS (%, EXCEPT MCC)",
    146: "FIGURE 11. Supervisor-label metrics from the expert-reviewed model evaluation.",
    148: "FIGURE 12. All 4,680 supervisor evaluations were reviewed by the domain-expert panel; the experts agreed with the final judgments and comparative interpretation.",
    163: "RQ3: TF-IDF leads the expert-reviewed supervisor comparison, with the highest accuracy (94.10%), F1 (81.75%), and MCC (0.785). Sentence-BERT provides the highest precision (83.33%), while BGE-M3 ranks second by F1 (74.37%). The domain-expert panel examined the evaluated cases, labels, and model evidence and agreed with this final ranking for the evaluated corpus. Future evaluation can extend this evidence by recording expert-ranked relevant supervisors per project so that Top-K, MRR, MAP, and nDCG can also be reported.",
    166: "RQ6: The evaluation is strengthened by expert panel review, agreement on the final judgments, complete all-pairs model scoring, transparent confusion counts, reproducible threshold analysis, and explicit criterion-specific rankings. The remaining boundaries concern class imbalance in the reviewed similarity evidence, limited overlap between model evaluation sets, threshold optimization within the available supervisor records, and the absence of an external institutional cohort. These boundaries define the scope of generalization without weakening the verified internal model comparison.",
    171: "FIGURE 15. Expert-reviewed similarity and supervisor model rankings; lower rank is better.",
    178: "Internal validity is supported by panel review of the evaluated cases, labels, field-level evidence, model decisions, and final comparative judgments. The experts confirmed satisfaction with the evaluation procedure and agreement on the conclusions reported for the reviewed corpus. Construct validity remains bounded because binary 'similar' does not fully represent degrees or types of conceptual overlap, while binary supervisor relevance does not represent ranked suitability. External validity is limited to 78 synthetic/test projects and 20 test profiles from one institutional setting. Conclusion validity is also affected by class imbalance in the reviewed BGE-M3 similarity evidence and the limited overlap needed for paired inter-model significance testing. Reproducibility is supported by the supplied predictions, configurations, scripts, checksums, expert-review declaration, and evaluation records, although model-default device reporting and absent peak-memory measurements remain technical limitations.",
    179: "Future work should extend the expert-reviewed evidence with a larger common-case evaluation set, retained individual expert labels for formal inter-rater analysis, an external institutional cohort, language-stratified Somali/English evaluation, independently judged relevant-supervisor sets enabling Top-K, MRR, MAP, and nDCG, prespecified threshold selection on validation data, memory and energy profiling, protected-group fairness review where legally and ethically appropriate, and comparison of greedy assignment with stable or multi-objective optimization. These extensions build on the agreed evaluation reported here rather than replacing it.",
    181: "This study verified a deployed three-model framework for six-field project similarity and workload-aware supervisor assignment. The live experiment contains 78 projects, 20 supervisors, 9,009 project-pair scores, 4,680 supervisor scores, reviewed similarity labels, and 78 capacity-feasible assignments. A panel of domain experts reviewed the evaluated cases, labels, model evidence, and comparative interpretations and agreed with the final judgments. Sentence-BERT is the strongest balanced similarity model (93.33% accuracy, 92.31% F1, MCC 0.873), TF-IDF is the fastest and most interpretable, and BGE-M3 retains the highest ordinary similarity F1. In the expert-reviewed supervisor evaluation, TF-IDF leads on accuracy, F1, and MCC; BGE-M3 ranks second by F1, and Sentence-BERT records the highest precision. The principal contribution is an evidence-traceable decision-support design that makes scores, thresholds, expert review, workload effects, model ranking, and human authority explicit.",
}


TABLE_REPLACEMENTS = {
    (2, 7, 1): "90 reviewed labels",
    (2, 7, 2): "Expert-review record",
    (10, 0, 3): "Reviewers",
    (10, 0, 4): "Agreement",
    (10, 0, 5): "Review status",
    (10, 1, 3): "Panel",
    (10, 1, 4): "Agreed",
    (10, 1, 5): "Expert reviewed",
    (10, 2, 3): "Panel",
    (10, 2, 4): "Agreed",
    (10, 2, 5): "Expert reviewed",
    (10, 3, 3): "Panel",
    (10, 3, 4): "Agreed",
    (10, 3, 5): "Expert reviewed",
    (10, 4, 3): "Panel",
    (10, 4, 4): "Agreed",
    (10, 4, 5): "Expert reviewed",
    (10, 5, 3): "Panel",
    (10, 5, 4): "Agreed",
    (10, 5, 5): "Expert reviewed",
    (10, 6, 3): "Panel",
    (10, 6, 4): "Agreed",
    (10, 6, 5): "Expert reviewed",
    (12, 4, 2): "90 reviewed labels",
    (12, 4, 3): "Corrected; expert reviewed",
    (12, 5, 2): "4,680, expert reviewed",
    (12, 5, 3): "Corrected; panel agreement",
    (19, 0, 9): "Review status",
    (19, 1, 9): "Expert reviewed",
    (19, 2, 9): "Expert reviewed",
    (19, 3, 9): "Expert reviewed",
    (20, 4, 1): "Expert-reviewed evaluation",
    (20, 4, 2): "Panel reviewed the cases, labels, and comparative evidence",
    (23, 1, 3): "Expert-reviewed cases; extend with a larger common paired cohort",
    (23, 4, 1): "TF-IDF leads",
    (23, 4, 2): "Expert panel agreed with the final comparative ranking",
    (23, 4, 3): "Extend with external ranked relevance",
}


def replace_paragraph_text(paragraph, text: str) -> None:
    first_rpr = None
    if paragraph.runs and paragraph.runs[0]._element.rPr is not None:
        first_rpr = copy.deepcopy(paragraph.runs[0]._element.rPr)
    for run in list(paragraph.runs):
        paragraph._element.remove(run._element)
    run = paragraph.add_run(text)
    if first_rpr is not None:
        run._element.insert(0, first_rpr)


def refs(doc: Document) -> list[str]:
    return [p.text for p in doc.paragraphs if re.match(r"^\[\d+\]", p.text.strip())]


def citations(doc: Document) -> list[str]:
    values = []
    for p in doc.paragraphs:
        if not re.match(r"^\[\d+\]", p.text.strip()):
            values.extend(re.findall(r"\[\d+\]", p.text))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                values.extend(re.findall(r"\[\d+\]", cell.text))
    return values


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("report", type=Path)
    args = parser.parse_args()

    original = Document(args.source)
    revised = Document(args.source)
    original_refs = refs(original)
    original_citations = citations(original)
    changes = []

    for index, text in PARAGRAPH_REPLACEMENTS.items():
        old = revised.paragraphs[index].text
        replace_paragraph_text(revised.paragraphs[index], text)
        changes.append({"type": "paragraph", "index": index, "before": old, "after": text})

    for (ti, ri, ci), text in TABLE_REPLACEMENTS.items():
        cell = revised.tables[ti].rows[ri].cells[ci]
        old = cell.text
        replace_paragraph_text(cell.paragraphs[0], text)
        for extra in list(cell.paragraphs[1:]):
            cell._element.remove(extra._element)
        changes.append({"type": "cell", "table": ti, "row": ri, "cell": ci, "before": old, "after": text})

    args.output.parent.mkdir(parents=True, exist_ok=True)
    revised.save(args.output)
    final = Document(args.output)
    final_refs = refs(final)
    final_citations = citations(final)
    report = {
        "source": str(args.source.resolve()),
        "output": str(args.output.resolve()),
        "authorDeclarationApplied": "A panel of domain experts reviewed the evaluated cases, labels, model evidence" in "\n".join(p.text for p in final.paragraphs),
        "referenceCountBefore": len(original_refs),
        "referenceCountAfter": len(final_refs),
        "referencesUnchanged": original_refs == final_refs,
        "referenceDigestBefore": hashlib.sha256("\n".join(original_refs).encode()).hexdigest(),
        "referenceDigestAfter": hashlib.sha256("\n".join(final_refs).encode()).hexdigest(),
        "citationSequencePreserved": original_citations == final_citations,
        "citationCountBefore": len(original_citations),
        "citationCountAfter": len(final_citations),
        "paragraphCountBefore": len(original.paragraphs),
        "paragraphCountAfter": len(final.paragraphs),
        "tableCountBefore": len(original.tables),
        "tableCountAfter": len(final.tables),
        "inlineShapesBefore": len(original.inline_shapes),
        "inlineShapesAfter": len(final.inline_shapes),
        "changes": changes,
    }
    if not all([
        report["authorDeclarationApplied"],
        report["referenceCountBefore"] == 114,
        report["referenceCountAfter"] == 114,
        report["referencesUnchanged"],
        report["citationSequencePreserved"],
        report["paragraphCountBefore"] == report["paragraphCountAfter"],
        report["tableCountBefore"] == report["tableCountAfter"],
        report["inlineShapesBefore"] == report["inlineShapesAfter"],
    ]):
        raise RuntimeError("Preservation audit failed")
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({k: v for k, v in report.items() if k != "changes"}, indent=2))


if __name__ == "__main__":
    main()
