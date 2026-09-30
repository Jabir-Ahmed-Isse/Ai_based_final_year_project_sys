from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "outputs" / "journal-package-20260724"


def configure(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    for style_name, size, color in [
        ("Title", 18, "17365D"),
        ("Heading 1", 14, "17365D"),
        ("Heading 2", 12, "2F5597"),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)


def add_inline_markdown(paragraph, text: str) -> None:
    position = 0
    for match in re.finditer(r"\*\*(.+?)\*\*|\*(.+?)\*", text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        run = paragraph.add_run(match.group(1) or match.group(2))
        run.bold = bool(match.group(1))
        run.italic = bool(match.group(2))
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def markdown_to_docx(source: Path, target: Path) -> None:
    document = Document()
    configure(document)
    for raw_line in source.read_text(encoding="utf-8").splitlines():
        line = raw_line.rstrip()
        if not line:
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(line[2:].strip())
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            document.add_heading(line[4:].strip(), level=2)
        elif line.startswith("- ["):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_markdown(paragraph, line[2:])
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_markdown(paragraph, line[2:])
        else:
            paragraph = document.add_paragraph()
            add_inline_markdown(paragraph, line.replace("  ", ""))
    document.save(target)


markdown_to_docx(PACKAGE / "cover-letter.md", PACKAGE / "cover-letter.docx")
markdown_to_docx(
    PACKAGE / "reviewer-response-template.md",
    PACKAGE / "reviewer-response-template.docx",
)

print(PACKAGE / "cover-letter.docx")
print(PACKAGE / "reviewer-response-template.docx")
