from __future__ import annotations

import argparse
import copy
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


TABLE_TITLE = "TABLE I. COMPARATIVE SCOPE OF FOUR RELATED STUDIES AND THE PRESENT RESEARCH"
INTRO_TEXT = (
    "Table I compares the present work with four closely related studies selected "
    "because they represent publication-based supervisor retrieval, content-based "
    "recommendation, preference-and-capacity optimization, and workload-balanced "
    "allocation. 'Yes' denotes a capability explicitly reported in the cited study; "
    "'NR' means that the capability was not reported in that study's stated "
    "method or evaluation. The comparison shows that earlier studies provide strong "
    "solutions to individual parts of the problem, whereas the present research joins "
    "proposal-overlap screening, comparative semantic-model evaluation, supervisor "
    "recommendation, capacity-aware assignment, and expert-reviewed evidence in one "
    "auditable workflow [35], [37], [38], [40]."
)

HEADERS = [
    "Study and reported method",
    "Proposal-overlap\nscreening",
    "Multiple-model\ncomparison",
    "Supervisor\nrecommendation",
    "Capacity/workload\nallocation",
    "Expert-reviewed\nintegrated workflow",
]

ROWS = [
    ["Damayanti et al. [35]: TF-IDF, K-means and cosine similarity", "NR", "NR", "Yes", "NR", "NR"],
    ["Sanchez-Anguix et al. [37]: multi-objective genetic allocation", "NR", "NR", "NR", "Yes", "NR"],
    ["Wijanto et al. [38]: publication profiles, vector-space retrieval and cosine similarity", "NR", "NR", "Yes", "NR", "NR"],
    ["Ramotsisi et al. [40]: preference-based optimization model", "NR", "NR", "NR", "Yes", "NR"],
    ["Present study: TF-IDF, Sentence-BERT and BGE-M3", "Yes", "Yes", "Yes", "Yes", "Yes"],
]

REFERENCE_UPDATES = {
    35: "[35] Damayanti, A., Purwani, F., & Kadafi, M. (2025). A content-based thesis supervisor recommendation system based on research interest clustering and cosine similarity. JUSIFO (Jurnal Sistem Informasi), 11(2), 111-120. https://doi.org/10.19109/jusifo.v11i2.27605",
    37: "[37] Sanchez-Anguix, V., Chalumuri, R., Aydogan, R., & Julian, V. (2019). A near Pareto optimal approach to student-supervisor allocation with two-sided preferences and workload balance. Applied Soft Computing, 76, 1-15. https://doi.org/10.1016/j.asoc.2018.11.049",
    38: "[38] Wijanto, M. C., Rachmadiany, R., & Karnalim, O. (2020). Thesis supervisor recommendation with representative content and information retrieval. Journal of Information Systems Engineering and Business Intelligence, 6(2), 143-150. https://doi.org/10.20473/jisebi.6.2.143-150",
    40: "[40] Ramotsisi, J., Kgomotso, M., Seboni, L., & Salahi, M. (2022). An optimization model for the student-to-project supervisor assignment problem: The case of an engineering department. Journal of Optimization, 2022, 9415210. https://doi.org/10.1155/2022/9415210",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=45, start=55, bottom=45, end=55) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "60")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def style_table(table) -> None:
    widths = [1300, 560, 660, 610, 660, 610]
    set_table_geometry(table, widths)
    table.autofit = False
    table.style = "Table Grid"
    for row_index, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        if row_index == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, "17365D")
            elif row_index == len(table.rows) - 1:
                set_cell_shading(cell, "D9EAF7")
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F2F4F7")
            else:
                set_cell_shading(cell, "FFFFFF")
            for paragraph in cell.paragraphs:
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT
                    if col_index == 0
                    else WD_ALIGN_PARAGRAPH.CENTER
                )
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(6.5)
                    if row_index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                    elif row_index == len(table.rows) - 1:
                        run.bold = True


def add_caption(doc: Document) -> object:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(TABLE_TITLE)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor(23, 54, 93)
    return paragraph


def add_intro(doc: Document) -> object:
    paragraph = doc.add_paragraph(INTRO_TEXT)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3)
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(8.5)
    return paragraph


def renumber_table_captions(doc: Document) -> int:
    count = 0
    pattern = re.compile(r"^(TABLE)\s+([IVXLCDM]+)(\..*)$", re.IGNORECASE)
    roman = {
        1: "I", 2: "II", 3: "III", 4: "IV", 5: "V", 6: "VI", 7: "VII",
        8: "VIII", 9: "IX", 10: "X", 11: "XI", 12: "XII", 13: "XIII",
        14: "XIV", 15: "XV", 16: "XVI", 17: "XVII", 18: "XVIII",
    }
    for paragraph in doc.paragraphs:
        match = pattern.match(paragraph.text.strip())
        if not match:
            continue
        count += 1
        new_text = f"TABLE {roman[count]}{match.group(3)}"
        if len(paragraph.runs) == 1:
            paragraph.runs[0].text = new_text
        else:
            paragraph.text = new_text
    return count


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("report", type=Path)
    args = parser.parse_args()

    source = Document(args.source)
    source_refs = [p.text for p in source.paragraphs if re.match(r"^\[\d+\]", p.text.strip())]
    doc = Document(args.source)
    anchor = doc.paragraphs[27]

    references_updated = 0
    for paragraph in doc.paragraphs:
        match = re.match(r"^\[(\d+)\]", paragraph.text.strip())
        if not match:
            continue
        number = int(match.group(1))
        if number in REFERENCE_UPDATES:
            if len(paragraph.runs) == 1:
                paragraph.runs[0].text = REFERENCE_UPDATES[number]
            else:
                paragraph.text = REFERENCE_UPDATES[number]
            references_updated += 1

    intro = add_intro(doc)
    caption = add_caption(doc)
    table = doc.add_table(rows=1, cols=len(HEADERS))
    for cell, value in zip(table.rows[0].cells, HEADERS):
        cell.text = value
    for values in ROWS:
        row = table.add_row()
        for cell, value in zip(row.cells, values):
            cell.text = value
    style_table(table)

    anchor._p.addnext(intro._p)
    intro._p.addnext(caption._p)
    caption._p.addnext(table._tbl)

    caption_count = renumber_table_captions(doc)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)

    final = Document(args.output)
    final_refs = [p.text for p in final.paragraphs if re.match(r"^\[\d+\]", p.text.strip())]
    report = {
        "comparisonTableInserted": TABLE_TITLE in "\n".join(p.text for p in final.paragraphs),
        "comparisonRows": 5,
        "comparisonStudies": [35, 37, 38, 40],
        "tableCaptionCount": caption_count,
        "tableCount": len(final.tables),
        "referenceCount": len(final_refs),
        "referencesUpdated": references_updated,
        "referenceNumbersPreserved": [re.match(r"^\[(\d+)\]", r.strip()).group(1) for r in final_refs]
        == [re.match(r"^\[(\d+)\]", r.strip()).group(1) for r in source_refs],
        "paragraphsAdded": len(final.paragraphs) - len(source.paragraphs),
        "inlineShapesPreserved": len(final.inline_shapes) == len(source.inline_shapes),
    }
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
