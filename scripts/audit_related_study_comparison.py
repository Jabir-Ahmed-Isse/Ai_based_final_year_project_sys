from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("pdf", type=Path)
    parser.add_argument("report", type=Path)
    args = parser.parse_args()

    doc = Document(args.docx)
    text = "\n".join(p.text for p in doc.paragraphs)
    all_text = text + "\n" + "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells
    )
    refs = [p.text.strip() for p in doc.paragraphs if re.match(r"^\[\d+\]", p.text.strip())]
    numbers = [int(re.match(r"^\[(\d+)\]", r).group(1)) for r in refs]
    captions = [p.text for p in doc.paragraphs if re.match(r"^TABLE\s+[IVXLCDM]+\.", p.text.strip())]
    expected_captions = [
        "I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X",
        "XI", "XII", "XIII", "XIV", "XV", "XVI",
    ]
    caption_numbers = [re.match(r"^TABLE\s+([IVXLCDM]+)\.", p.strip()).group(1) for p in captions]
    doi_checks = {
        35: "10.19109/jusifo.v11i2.27605",
        37: "10.1016/j.asoc.2018.11.049",
        38: "10.20473/jisebi.6.2.143-150",
        40: "10.1155/2022/9415210",
    }
    doi_present = {
        number: any(r.startswith(f"[{number}]") and doi in r for r in refs)
        for number, doi in doi_checks.items()
    }
    comparison_tables = [
        t for t in doc.tables
        if t.rows and t.rows[0].cells[0].text.strip() == "Study and reported method"
    ]
    pdf = PdfReader(str(args.pdf))
    blank_pages = [i + 1 for i, page in enumerate(pdf.pages) if not (page.extract_text() or "").strip()]
    with ZipFile(args.docx) as archive:
        crc = archive.testzip()
        media = len([n for n in archive.namelist() if n.startswith("word/media/")])

    output = {
        "sha256": hashlib.sha256(args.docx.read_bytes()).hexdigest(),
        "zipCrcFailure": crc,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "inlineShapes": len(doc.inline_shapes),
        "mediaParts": media,
        "comparisonTableCount": len(comparison_tables),
        "comparisonRowsIncludingHeader": len(comparison_tables[0].rows) if comparison_tables else 0,
        "comparisonColumns": len(comparison_tables[0].columns) if comparison_tables else 0,
        "fourCitedStudiesPresent": all(f"[{n}]" in all_text for n in [35, 37, 38, 40]),
        "presentStudyRowPresent": "Present study: TF-IDF, Sentence-BERT and BGE-M3" in all_text,
        "references": len(refs),
        "referencesSequential": numbers == list(range(1, 115)),
        "verifiedDoisPresent": doi_present,
        "tableCaptionsSequential": caption_numbers == expected_captions,
        "pdfPages": len(pdf.pages),
        "blankPdfPages": blank_pages,
    }
    output["passed"] = all([
        crc is None,
        len(doc.tables) == 25,
        len(doc.inline_shapes) == 15,
        media == 15,
        len(comparison_tables) == 1,
        len(comparison_tables[0].rows) == 6,
        len(comparison_tables[0].columns) == 6,
        output["fourCitedStudiesPresent"],
        output["presentStudyRowPresent"],
        len(refs) == 114,
        output["referencesSequential"],
        all(doi_present.values()),
        output["tableCaptionsSequential"],
        len(pdf.pages) == 21,
        not blank_pages,
    ])
    args.report.write_text(json.dumps(output, indent=2), encoding="utf-8")
    print(json.dumps(output, indent=2))
    if not output["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
