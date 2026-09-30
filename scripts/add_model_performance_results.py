from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from build_research_report import (
    MUTED,
    add_heading,
    add_paragraph,
    add_table,
    set_font,
)


ABSTRACT_REPLACEMENT = (
    "Final-year projects are often managed through fragmented processes for "
    "originality screening, supervisor selection, communication and monitoring. "
    "This study evaluates an implemented lifecycle platform integrating six-field "
    "project similarity with workload-aware supervisor recommendation. Using 78 "
    "Computer Science and IT project records and 20 structured test supervisor "
    "profiles, the reproducible experiment generated all 3,003 unique project pairs "
    "per model (9,009 records) and all 1,560 project-supervisor combinations per model "
    "(4,680 records) for TF-IDF, multilingual Sentence-BERT and BGE-M3. Mean "
    "unlabelled similarity scores were 18.93%, 47.11% and 63.07%, respectively. "
    "BGE-M3 therefore ranked first by raw score level, Sentence-BERT second and "
    "TF-IDF third; this ordering measures calibration and alert propensity, not "
    "accuracy. Project-pair Spearman correlations ranged from 0.889 to 0.954, with "
    "Sentence-BERT and BGE-M3 showing the strongest agreement (ρₛ = 0.954) and the "
    "same first-choice supervisor for 71.79% of projects. The evidence-bounded "
    "deployment ranking was TF-IDF first for speed and transparency, Sentence-BERT "
    "second as the balanced semantic option and BGE-M3 third as an offline research "
    "candidate because of its 818.871-s initial load and 1,024-dimensional "
    "representation. A capacity-aware ensemble assigned all 78 projects across all "
    "20 supervisors, with mean load 3.90 (SD 1.09), range 2-6 and no capacity "
    "violation. A metric-complete model matrix reports expert-ground-truth accuracy, "
    "precision, recall, macro F1 and MCC as NE (not estimable) for TF-IDF, "
    "Sentence-BERT and BGE-M3 because the frozen experiment contains zero "
    "expert-labelled pairs; the cells are shown explicitly rather than omitted. The "
    "thesis-reported 307-comparison dashboard values (96.42% accuracy, 95.65% "
    "precision, 99.00% recall and 97.29% F1) are retained separately as historical, "
    "model-unattributed evidence because the pair labels, predictions and confusion "
    "matrix are unavailable. The contribution is a complete operational comparison, "
    "explainable multi-field evidence, an annotation workflow and a responsible "
    "validation protocol for human-controlled academic decision support."
)


REPORT_ABSTRACT_REPLACEMENT = (
    "This report documents a complete local research experiment using 78 imported "
    "Computer Science and IT final-year projects, 20 structured supervisor profiles, "
    "and three similarity model families: TF-IDF, multilingual Sentence-BERT, and "
    "BGE-M3. The system generated and stored all 3,003 unique project pairs per model "
    "(9,009 model-pair records) and all 1,560 project-supervisor combinations per "
    "model (4,680 matching records). A capacity-aware ensemble procedure produced 78 "
    "provisional assignments using all 20 supervisors with no capacity violations. "
    "The frozen evidence contains zero independent expert-consensus pair labels. A "
    "metric-complete matrix therefore shows accuracy, precision, recall, macro F1 and "
    "MCC as NE (not estimable) for every completed model, while reporting measured "
    "coverage, score distributions, correlations, field scores, timings and workload "
    "results in full. The thesis's historical 307-comparison dashboard values "
    "(96.42% accuracy, 95.65% precision, 99.00% recall and 97.29% F1) are recorded "
    "separately because no model identifier, pair-level reference labels, predictions "
    "or confusion matrix are available to reproduce or assign them. A separate "
    "600-pair synthetic construction benchmark is supplied only to prepare the "
    "validation workflow; it is not presented as human-validated evidence."
)


def find_paragraph(document: Document, predicate, description: str):
    matches = [paragraph for paragraph in document.paragraphs if predicate(paragraph)]
    if len(matches) != 1:
        raise ValueError(
            f"Expected exactly one paragraph for {description}; found {len(matches)}"
        )
    return matches[0]


def replace_paragraph_text(paragraph, text: str) -> None:
    paragraph.text = text


def document_text(document: Document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def add_caption_paragraph(document: Document, text: str):
    paragraph = document.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_font(run, size=9, color=MUTED, italic=True)
    return paragraph


def performance_rows() -> list[list[str]]:
    return [
        ["Scored project-pair coverage", "3,003/3,003", "3,003/3,003", "3,003/3,003", "Complete and equal"],
        ["Independent expert-labelled pairs", "0", "0", "0", "No gold-standard test set"],
        ["Accuracy", "NE", "NE", "NE", "Requires independent reference labels"],
        ["Precision", "NE", "NE", "NE", "Requires true- and false-positive counts"],
        ["Recall", "NE", "NE", "NE", "Requires true-positive and false-negative counts"],
        ["Macro F1", "NE", "NE", "NE", "Requires classwise precision and recall"],
        ["MCC", "NE", "NE", "NE", "Requires a valid confusion matrix"],
        ["Mean similarity score", "18.93% (#3)", "47.11% (#2)", "63.07% (#1)", "Raw score level; not accuracy"],
        ["Median similarity score", "5.26%", "40.94%", "58.55%", "Model-specific calibration"],
        ["Low / medium / high alerts", "2,268 / 673 / 62", "1,390 / 1,210 / 403", "0 / 1,949 / 1,054", "Common 40% and 70% cut points"],
        ["Mean pairwise Spearman ρₛ", "0.8971 (#3)", "0.9215 (#2)", "0.9293 (#1)", "Cross-model rank agreement only"],
        ["Mean top-supervisor agreement", "62.82% (#3)", "68.59% (#1)", "66.03% (#2)", "Pairwise choice agreement; not relevance"],
        ["Initial model load", "0.000 s (#1)", "271.849 s (#2)", "818.871 s (#3)", "Lower is operationally better"],
        ["Stored pair computation", "295.284 ms", "281.698 s", "18.874 ms cached", "BGE-M3 value excludes cold loading"],
        ["Supervisor computation", "0.393 s (#1)", "10.786 s (#2)", "98.680 s (#3)", "All 1,560 candidates/model"],
        ["Representation dimension", "Vocabulary-dependent", "384", "1,024", "Size is not predictive quality"],
        ["Raw-score rank", "#3", "#2", "#1", "Calibration and alert propensity"],
        ["Agreement rank", "#3", "#2", "#1", "Mean pairwise Spearman only"],
        ["Operational deployment rank", "#1", "#2", "#3", "Current speed, transparency and cost evidence"],
        ["Predictive-quality rank", "NE", "NE", "NE", "Withheld until held-out expert evaluation"],
    ]


def insert_performance_section(
    document: Document,
    anchor,
    heading_text: str,
    heading_level: int,
    caption_text: str,
) -> None:
    heading = add_heading(document, heading_text, level=heading_level)
    lead = add_paragraph(
        document,
        (
            "Primary performance finding. The completed experiment provides full "
            "model outputs but zero independent expert-consensus labels. Consequently, "
            "the table reports every requested classification metric for every model "
            "and marks each unavailable gold-standard value as NE rather than leaving "
            "it blank or substituting a model-derived pseudo-label."
        ),
        bold_lead="Primary performance finding.",
    )
    caption = add_caption_paragraph(document, caption_text)
    table = add_table(
        document,
        ["Metric", "TF-IDF", "Sentence-BERT", "BGE-M3", "Interpretation"],
        performance_rows(),
        [1600, 1370, 1500, 1370, 3520],
        font_size=7.2,
    )
    spacer = document.paragraphs[-1]
    note = add_paragraph(
        document,
        (
            "NE means not estimable from the frozen evidence. The values in the score, "
            "agreement and timing rows are measured model behaviours; they are not "
            "substitutes for accuracy, precision, recall, F1 or MCC. A predictive "
            "winner can be ranked only after blinded labels are joined to the stored "
            "pair identifiers and evaluated on a held-out test split."
        ),
    )
    legacy = add_paragraph(
        document,
        (
            "Historical thesis dashboard result (not assigned to a model). The thesis "
            "reports n = 307 comparisons, accuracy = 96.42%, precision = 95.65%, "
            "recall = 99.00% and F1 = 97.29%. The inspected evidence does not contain "
            "the pair-level reference labels, prediction vector, model identifier or "
            "confusion matrix required to reproduce or attribute those values. They "
            "are retained for traceability but excluded from the TF-IDF, "
            "Sentence-BERT and BGE-M3 performance ranking."
        ),
        bold_lead="Historical thesis dashboard result (not assigned to a model).",
    )
    elements = [
        heading._p,
        lead._p,
        caption._p,
        table._tbl,
        spacer._p,
        note._p,
        legacy._p,
    ]
    for element in reversed(elements):
        anchor._p.addnext(element)


def renumber_journal_table_captions(document: Document) -> None:
    number_map = {8: 9, 9: 10, 10: 11, 11: 12}
    for paragraph in document.paragraphs:
        match = re.match(r"^Table (\d+)\.(.*)$", paragraph.text)
        if not match:
            continue
        old_number = int(match.group(1))
        if old_number in number_map:
            paragraph.text = f"Table {number_map[old_number]}.{match.group(2)}"


def add_journal_highlight(document: Document) -> None:
    anchor = find_paragraph(
        document,
        lambda paragraph: paragraph.text
        == "Ground-truth absence is explicitly separated from descriptive model behaviour.",
        "ground-truth highlight",
    )
    highlight = document.add_paragraph(style="List Bullet")
    highlight.add_run(
        "Accuracy, precision, recall, F1 and MCC are shown explicitly for every "
        "completed model; all are NE until expert labels exist."
    )
    anchor._p.addnext(highlight._p)


def update_journal(source: Path, destination: Path) -> None:
    document = Document(source)
    original_media_count = len(document.inline_shapes)
    original_table_count = len(document.tables)

    abstract = find_paragraph(
        document,
        lambda paragraph: paragraph.text.startswith(
            "Final-year projects are often managed through fragmented processes"
        ),
        "journal abstract",
    )
    replace_paragraph_text(abstract, ABSTRACT_REPLACEMENT)
    add_journal_highlight(document)

    comparison_intro = find_paragraph(
        document,
        lambda paragraph: paragraph.text.startswith(
            "The comparative report was integrated into the article"
        ),
        "comparative evaluation introduction",
    )
    replace_paragraph_text(
        comparison_intro,
        (
            "The comparative report was integrated into the article by separating six "
            "questions that would otherwise be conflated: Which model produces the "
            "highest raw scores? Which model separates this corpus most strongly? "
            "Which model agrees most closely with the others? Which model is least "
            "expensive to operate? Which model is the most appropriate staged-"
            "deployment choice? Which model has the best expert-validated predictive "
            "performance? A rank is reported only within its named dimension. The "
            "last question is answered explicitly as not estimable for every model "
            "because no expert ground truth exists."
        ),
    )

    renumber_journal_table_captions(document)
    bge_anchor = find_paragraph(
        document,
        lambda paragraph: paragraph.text.startswith(
            "BGE-M3 ranked first by raw score level"
        ),
        "BGE-M3 comparison paragraph",
    )
    insert_performance_section(
        document,
        bge_anchor,
        "4.4.4 Metric-complete model-performance results",
        3,
        (
            "Table 8. Complete per-model performance matrix. NE = not estimable from "
            "the current zero-label evidence."
        ),
    )

    statistical_boundary = find_paragraph(
        document,
        lambda paragraph: paragraph.text.startswith(
            "No confusion matrix, ROC curve, precision-recall curve"
        ),
        "statistical boundary paragraph",
    )
    replace_paragraph_text(
        statistical_boundary,
        (
            "Table 8 explicitly reports accuracy, precision, recall, macro F1 and MCC "
            "as NE for TF-IDF, Sentence-BERT and BGE-M3 because the labelled evaluation "
            "count is zero for each model. No confusion matrix, ROC curve, precision-"
            "recall curve, F1 threshold curve, McNemar test or prediction-error "
            "significance test is reported because no expert outcome exists. Likewise, "
            "Top-1/3/5 accuracy, MRR, MAP and NDCG cannot be calculated for supervisor "
            "ranking without relevance grades or accepted-assignment alternatives. "
            "The valid current analyses are raw score distributions, field scores, "
            "cross-model association, timing, coverage, feasibility and disagreement "
            "review."
        ),
    )

    historical_boundary = find_paragraph(
        document,
        lambda paragraph: paragraph.text.startswith(
            "The thesis reports 307 comparisons with 96.42% accuracy"
        ),
        "historical thesis metric paragraph",
    )
    replace_paragraph_text(
        historical_boundary,
        (
            "The thesis reports 307 comparisons with 96.42% accuracy, 95.65% "
            "precision, 99.00% recall and 97.29% F1. These historical values are shown "
            "separately beneath Table 8, but the inspected workspace does not contain "
            "the model identifier, pair-level labels, prediction vector or confusion "
            "matrix required to reproduce or assign them to TF-IDF, Sentence-BERT or "
            "BGE-M3. They are therefore excluded from the primary model ranking. This "
            "decision resolves the conflict in favor of auditable database evidence."
        ),
    )

    conclusion_opening = find_paragraph(
        document,
        lambda paragraph: paragraph.text.startswith(
            "The implemented platform operationalizes proposal archiving"
        ),
        "journal conclusion opening",
    )
    replace_paragraph_text(
        conclusion_opening,
        (
            "The implemented platform operationalizes proposal archiving, multi-field "
            "similarity screening, supervisor recommendation, capacity-aware "
            "provisional assignment and lifecycle administration. TF-IDF, "
            "Sentence-BERT and BGE-M3 successfully produced complete score sets. The "
            "metric-complete comparison now shows labelled n = 0 and accuracy, "
            "precision, recall, macro F1 and MCC = NE for each model; RQ1-RQ3 therefore "
            "cannot be answered in terms of comparative predictive quality until "
            "expert labels exist. The observable answer is that score distributions, "
            "representation sizes, agreement and run costs differ substantially. RQ4 "
            "is answered positively for feasibility: all 78 projects were assigned "
            "across 20 test supervisors with no capacity violation. RQ5 identifies the "
            "decisive next requirement: blinded expert annotation of pair similarity "
            "and multi-relevant supervisor suitability, followed by validation-only "
            "threshold/weight selection and held-out evaluation."
        ),
    )

    conclusion_rank = find_paragraph(
        document,
        lambda paragraph: paragraph.text.startswith(
            "The integrated comparison supports two deliberately different rankings"
        ),
        "journal conclusion ranking paragraph",
    )
    replace_paragraph_text(
        conclusion_rank,
        (
            "The integrated comparison supports three deliberately different measured "
            "rankings and one withheld ranking. Raw unlabelled score level ranks "
            "BGE-M3 first (63.07%), Sentence-BERT second (47.11%) and TF-IDF third "
            "(18.93%); this ranking measures calibration and alert propensity only. "
            "Mean pairwise rank agreement places BGE-M3 first (ρₛ = 0.9293), "
            "Sentence-BERT second (ρₛ = 0.9215) and TF-IDF third (ρₛ = 0.8971). The "
            "evidence-bounded deployment ranking reverses the extremes: TF-IDF is "
            "first as the fastest and most transparent baseline, Sentence-BERT is "
            "second as the balanced semantic option and BGE-M3 is third as an offline "
            "research candidate. Predictive-quality rank is NE for all three models; "
            "no winner is declared until independent labels permit held-out accuracy, "
            "precision, recall, F1, MCC and ranking evaluation."
        ),
    )

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)

    reopened = Document(destination)
    if len(reopened.tables) != original_table_count + 1:
        raise ValueError(
            f"Expected {original_table_count + 1} tables; found {len(reopened.tables)}"
        )
    if len(reopened.inline_shapes) != original_media_count:
        raise ValueError(
            f"Expected {original_media_count} inline figures; found "
            f"{len(reopened.inline_shapes)}"
        )
    required_text = [
        "4.4.4 Metric-complete model-performance results",
        "Accuracy",
        "Precision",
        "Recall",
        "Macro F1",
        "MCC",
        "Historical thesis dashboard result",
        "Table 12. Reconciliation of high-impact claims",
    ]
    all_text = document_text(reopened)
    missing = [item for item in required_text if item not in all_text]
    if missing:
        raise ValueError(f"Missing required journal content: {missing}")


def update_comparative_report(source: Path, destination: Path) -> None:
    document = Document(source)
    original_table_count = len(document.tables)
    original_media_count = len(document.inline_shapes)

    abstract = find_paragraph(
        document,
        lambda paragraph: paragraph.text.startswith(
            "This report documents a complete local research experiment"
        ),
        "comparative report abstract",
    )
    replace_paragraph_text(abstract, REPORT_ABSTRACT_REPLACEMENT)

    model_comparison = find_paragraph(
        document,
        lambda paragraph: paragraph.text.startswith("No quality rank is assigned."),
        "comparative report model comparison",
    )
    replace_paragraph_text(
        model_comparison,
        (
            "No predictive-quality rank is assigned. A higher unlabelled mean score can "
            "reflect calibration rather than better discrimination. The matrix below "
            "shows accuracy, precision, recall, macro F1 and MCC explicitly for every "
            "model as NE because all three labelled sample counts are zero. Measured "
            "score, agreement, coverage, latency and operational ranks remain fully "
            "reported. The final quality decision can be populated only after a "
            "preregistered, expert-labelled held-out evaluation."
        ),
    )
    insert_performance_section(
        document,
        model_comparison,
        "20.1 Metric-complete model-performance results",
        2,
        (
            "Table 9. Complete per-model performance matrix. NE = not estimable from "
            "the current zero-label evidence."
        ),
    )

    limitations = find_paragraph(
        document,
        lambda paragraph: paragraph.text.startswith(
            "No human-validated accuracy or ranking quality is available."
        ),
        "comparative report limitations",
    )
    replace_paragraph_text(
        limitations,
        (
            "The metric-complete matrix reports labelled n = 0 and accuracy, precision, "
            "recall, macro F1 and MCC = NE for each completed model. The thesis's "
            "historical 307-case metrics cannot be attributed or reproduced from the "
            "available evidence. The synthetic benchmark is not expert ground truth. "
            "BGE-M3 uses dense embeddings only; sparse and multi-vector modes were not "
            "activated. Memory was observed operationally but not captured with a "
            "controlled profiler for every run. Domain-level evaluation awaits a "
            "validated project-domain taxonomy."
        ),
    )

    conclusion = find_paragraph(
        document,
        lambda paragraph: paragraph.text.startswith(
            "The implementation completed the full unlabelled experiment safely"
        ),
        "comparative report conclusion",
    )
    replace_paragraph_text(
        conclusion,
        (
            "The implementation completed the full unlabelled experiment safely and "
            "reproducibly: 78 projects, 9,009 pair scores, 4,680 supervisor scores and "
            "78 capacity-safe assignments. Every requested classification metric is "
            "now visible for every model: labelled n = 0 and accuracy, precision, "
            "recall, macro F1 and MCC = NE for TF-IDF, Sentence-BERT and BGE-M3. This "
            "is the central performance result of the current evidence, not an omitted "
            "analysis. The package also provides the annotation, consensus, export, "
            "visualization and audit foundations required for a valid second phase. "
            "All three models are operational and their behaviours differ materially, "
            "but predictive quality remains undetermined until expert ground truth is "
            "collected."
        ),
    )

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)

    reopened = Document(destination)
    if len(reopened.tables) != original_table_count + 1:
        raise ValueError(
            f"Expected {original_table_count + 1} report tables; "
            f"found {len(reopened.tables)}"
        )
    if len(reopened.inline_shapes) != original_media_count:
        raise ValueError(
            f"Expected {original_media_count} report figures; "
            f"found {len(reopened.inline_shapes)}"
        )
    all_text = document_text(reopened)
    if "20.1 Metric-complete model-performance results" not in all_text:
        raise ValueError("Comparative report metric-complete section was not inserted")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Add explicit per-model performance-metric results to journal files."
    )
    parser.add_argument("--journal-source", required=True, type=Path)
    parser.add_argument("--journal-output", required=True, type=Path)
    parser.add_argument("--report-source", required=True, type=Path)
    parser.add_argument("--report-output", required=True, type=Path)
    args = parser.parse_args()

    update_journal(args.journal_source, args.journal_output)
    update_comparative_report(args.report_source, args.report_output)
    print(args.journal_output)
    print(args.report_output)


if __name__ == "__main__":
    main()
