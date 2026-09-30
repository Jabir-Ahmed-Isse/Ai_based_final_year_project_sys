from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree
from pypdf import PdfReader


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("pdf", type=Path)
    parser.add_argument("revision_report", type=Path)
    parser.add_argument("screenshot_dir", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    doc = Document(args.docx)
    paragraph_text = [p.text for p in doc.paragraphs]
    full_text = "\n".join(paragraph_text)
    references = []
    in_references = False
    for text in paragraph_text:
        stripped = text.strip()
        if stripped == "REFERENCES":
            in_references = True
            continue
        if in_references:
            match = re.match(r"^\[(\d+)\]", stripped)
            if match:
                references.append(int(match.group(1)))

    appendix_captions = [
        text.strip()
        for text in paragraph_text
        if re.match(r"^FIGURE A(?:[1-9]|1\d|2[0-2])\.", text.strip(), re.I)
    ]

    table_reference_rows = []
    traceability = {}
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if cells and cells[0] == "References":
                table_reference_rows.append(cells)
            if len(cells) >= 3 and re.fullmatch(r"A(?:[1-9]|1\d|2[0-2])", cells[0]):
                traceability[cells[0]] = cells[2]

    with ZipFile(args.docx) as archive:
        bad_zip_part = archive.testzip()
        xml = etree.fromstring(archive.read("word/document.xml"))
        equation_count = len(xml.xpath(".//w:oMath | .//w:oMathPara", namespaces=NS))
        media_hashes = {
            Path(name).name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("word/media/")
        }

    revision = json.loads(args.revision_report.read_text(encoding="utf-8"))
    replacement_checks = {}
    for figure, item in revision["visualReplacements"].items():
        screenshot = args.screenshot_dir / item["sourceScreenshot"]
        replacement_checks[figure] = {
            "embeddedPart": item["embeddedPart"],
            "sourceScreenshot": item["sourceScreenshot"],
            "sourceExists": screenshot.exists(),
            "hashMatch": screenshot.exists()
            and media_hashes.get(item["embeddedPart"]) == sha256(screenshot),
            "traceabilityMatch": traceability.get(figure) == item["sourceScreenshot"],
        }

    pdf = PdfReader(str(args.pdf))
    page_text_lengths = [len((page.extract_text() or "").strip()) for page in pdf.pages]

    report = {
        "docx": str(args.docx.resolve()),
        "docxSha256": sha256(args.docx),
        "zipCrcFailure": bad_zip_part,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "inlineShapes": len(doc.inline_shapes),
        "equationElements": equation_count,
        "referenceEntries": len(references),
        "referencesSequential1To113": references == list(range(1, 114)),
        "appendixCaptions": len(appendix_captions),
        "appendixCaptionsSequentialA1ToA22": all(
            caption.upper().startswith(f"FIGURE A{i}.")
            for i, caption in enumerate(appendix_captions, start=1)
        ),
        "tableReferenceRows": table_reference_rows,
        "referenceCountCorrect": table_reference_rows == [
            ["References", "113", "113", "Confirmed; finalized list preserved"]
        ],
        "containsObsoleteReferenceCounts": any(
            row[:4] == ["References", "160", "210", "Expanded and reverified"]
            for row in table_reference_rows
        ),
        "replacementChecks": replacement_checks,
        "all15ReplacementHashesMatch": len(replacement_checks) == 15
        and all(item["hashMatch"] for item in replacement_checks.values()),
        "all15TraceabilityRowsMatch": len(replacement_checks) == 15
        and all(item["traceabilityMatch"] for item in replacement_checks.values()),
        "pdfPages": len(pdf.pages),
        "blankPdfPages": [i + 1 for i, length in enumerate(page_text_lengths) if length == 0],
        "methodologyIntegrityStatements": {
            "singleAdministrator": "submitted by one administrator" in full_text,
            "notMultiRaterConsensus": "not multi-rater consensus" in full_text,
            "supervisorSingleReviewer": "single-reviewer" in full_text,
            "dashboardConsensusCaptionClarified": (
                "Dashboard captions that used the word consensus are corrected in the manuscript"
                in full_text
            ),
        },
    }
    required = [
        report["zipCrcFailure"] is None,
        report["tables"] == 25,
        report["inlineShapes"] == 37,
        report["referenceEntries"] == 113,
        report["referencesSequential1To113"],
        report["appendixCaptions"] == 22,
        report["appendixCaptionsSequentialA1ToA22"],
        report["referenceCountCorrect"],
        not report["containsObsoleteReferenceCounts"],
        report["all15ReplacementHashesMatch"],
        report["all15TraceabilityRowsMatch"],
        report["pdfPages"] == 33,
        not report["blankPdfPages"],
        all(report["methodologyIntegrityStatements"].values()),
    ]
    report["passed"] = all(required)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(report, indent=2, ensure_ascii=False))
    if not report["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
