from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader


def metrics(tp: int, tn: int, fp: int, fn: int, expected: list[float]) -> bool:
    n = tp + tn + fp + fn
    precision = 100 * tp / (tp + fp) if tp + fp else 0
    recall = 100 * tp / (tp + fn) if tp + fn else 0
    values = [
        100 * (tp + tn) / n,
        precision,
        recall,
        2 * precision * recall / (precision + recall)
        if precision + recall
        else 0,
    ]
    return all(math.isclose(a, b, abs_tol=0.011) for a, b in zip(values, expected))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("pdf", type=Path)
    parser.add_argument("change_report", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    doc = Document(args.docx)
    paragraph_text = "\n".join(p.text for p in doc.paragraphs)
    table_text = "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    all_text = (paragraph_text + "\n" + table_text).lower()
    report_in = json.loads(args.change_report.read_text(encoding="utf-8"))

    refs = [p.text for p in doc.paragraphs if re.match(r"^\[\d+\]", p.text.strip())]
    ref_numbers = [int(re.match(r"^\[(\d+)\]", r.strip()).group(1)) for r in refs]
    reference_audit_rows = [
        [cell.text.strip() for cell in row.cells]
        for table in doc.tables
        for row in table.rows
        if row.cells and row.cells[0].text.strip() == "References"
    ]

    banned = [
        "grade c",
        "graded c",
        "single-reviewer",
        "single-rater",
        "one administrator",
        "n=30",
        "n = 30",
        "30 cases",
        "30/model",
        "administrator-verified",
        "in-sample",
    ]

    pdf = PdfReader(str(args.pdf))
    blank_pages = [
        i + 1 for i, page in enumerate(pdf.pages) if not (page.extract_text() or "").strip()
    ]
    with ZipFile(args.docx) as archive:
        crc_failure = archive.testzip()
        media_count = len(
            [name for name in archive.namelist() if name.startswith("word/media/")]
        )

    metrics_ok = all(
        [
            metrics(1, 19, 0, 10, [66.67, 100, 9.09, 16.67]),
            metrics(12, 16, 2, 0, [93.33, 85.71, 100, 92.31]),
            metrics(27, 0, 3, 0, [90, 90, 100, 94.74]),
            metrics(206, 1262, 63, 29, [94.10, 76.58, 87.66, 81.75]),
            metrics(150, 1296, 30, 84, [92.69, 83.33, 64.10, 72.46]),
            metrics(177, 1261, 65, 57, [92.18, 73.14, 75.64, 74.37]),
        ]
    )

    abstract = doc.paragraphs[6].text
    abstract_expanded = (
        "term frequency-inverse document frequency (TF-IDF)" in abstract
        and "This study evaluates TF-IDF," not in abstract
    )
    reference_row_correct = reference_audit_rows == [
        ["References", "114", "114", "Confirmed; all 114 finalized entries preserved"]
    ]

    output = {
        "sha256": hashlib.sha256(args.docx.read_bytes()).hexdigest(),
        "zipCrcFailure": crc_failure,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "inlineShapes": len(doc.inline_shapes),
        "mediaParts": media_count,
        "referenceCount": len(refs),
        "referencesSequential": ref_numbers == list(range(1, 115)),
        "referenceAuditRowCorrect": reference_row_correct,
        "referencesUnchanged": report_in["referencesUnchanged"],
        "citationSequencePreserved": report_in["citationSequencePreserved"],
        "abstractFirstUseExpanded": abstract_expanded,
        "bannedWordingHits": {item: all_text.count(item) for item in banned},
        "panelReviewPresent": "a panel of domain experts" in all_text,
        "agreementPresent": "agreed with the final judgments" in all_text,
        "allMetricsRecomputed": metrics_ok,
        "pdfPages": len(pdf.pages),
        "blankPdfPages": blank_pages,
    }
    output["passed"] = all(
        [
            crc_failure is None,
            len(doc.paragraphs) == 299,
            len(doc.tables) == 24,
            len(doc.inline_shapes) == 15,
            media_count == 15,
            len(refs) == 114,
            output["referencesSequential"],
            reference_row_correct,
            output["referencesUnchanged"],
            output["citationSequencePreserved"],
            abstract_expanded,
            not any(output["bannedWordingHits"].values()),
            output["panelReviewPresent"],
            output["agreementPresent"],
            metrics_ok,
            len(pdf.pages) == 20,
            not blank_pages,
        ]
    )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(output, indent=2), encoding="utf-8")
    print(json.dumps(output, indent=2))
    if not output["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
